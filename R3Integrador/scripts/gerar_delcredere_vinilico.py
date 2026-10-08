#!/usr/bin/env python3
"""Gera as seis importações Del Credere para os vinílicos da tabela 100."""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/home/fernando/Projetos/Work/ARCHOME/Planilhas/delcredere/100 TABELA DEL CREDERE SP ATUALIZADA.xlsx")
TEMPLATE = Path("/home/fernando/Projetos/Work/ARCHOME/Planilhas/IMPORTACAO_ERP_VINILICO_20260701_200706_CONTADOR.xlsx")
OUT_DIR = ROOT / "Saida" / "DELCREDERE_VINILICO_20260909"
TABLES = {"DEL5": 18, "DEL10": 19, "DEL15": 20, "DEL20": 21, "DEL25": 22, "DEL30": 23}


def code(value: object) -> str:
    return "" if value is None else str(value).replace("***", "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sale_price(factory: object) -> Decimal:
    # Regra já usada nas importações Del Credere: preço da faixa + IPI 0,65%
    # e taxa de cartão 4,71%, totalizando 5,36%.
    return (Decimal(str(factory)) * Decimal("1.0536")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def main() -> None:
    source_ws = load_workbook(SOURCE, read_only=True, data_only=True).active
    products: dict[str, dict[str, object]] = {}
    last_format = last_color = ""
    last_m2_box = last_weight_box = None
    for row in source_ws.iter_rows(min_row=276, max_row=291, values_only=True):
        reference = code(row[1])
        if not reference:
            continue
        if row[0] is not None:
            last_format = str(row[0]).replace(" ", "").upper()
        if row[4] is not None:
            last_color = str(row[4]).strip().lower()
        if row[10] is not None:
            last_m2_box = row[10]
        if row[15] is not None:
            last_weight_box = row[15]
        if not all((last_format, last_color, last_m2_box, last_weight_box)):
            raise ValueError(f"Dados comerciais ausentes para {reference}")
        products[reference] = {
            "modelo": last_format, "cor": last_color, "m2_caixa": last_m2_box,
            "peso_bruto_caixa": last_weight_box,
            "precos": {name: row[column - 1] for name, column in TABLES.items()},
        }
    if len(products) != 16:
        raise ValueError(f"Esperados 16 vinílicos; encontrados {len(products)}")

    OUT_DIR.mkdir(exist_ok=True)
    generated: dict[str, str] = {}
    for table, column in TABLES.items():
        book = load_workbook(TEMPLATE)
        ws = book.active
        legacy = {f"LVT9635000{i}": f"LVT9435000{i}" for i in range(1, 5)}
        references = [legacy.get(code(ws.cell(row, 2).value), code(ws.cell(row, 2).value)) for row in range(2, ws.max_row + 1)]
        if set(references) != set(products):
            raise ValueError(f"Referências do modelo não coincidem em {table}")
        for row, reference in zip(range(2, ws.max_row + 1), references):
            item = products[reference]
            factory = Decimal(str(item["precos"][table])).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            ws.cell(row, 2).value = reference
            ws.cell(row, 8).value = f"VILLAGRES {table.removeprefix('DEL')}"
            ws.cell(row, 10).value = item["modelo"]
            ws.cell(row, 12).value = item["cor"]
            ws.cell(row, 15).value = float(sale_price(factory))
            ws.cell(row, 16).value = float(factory)
            ws.cell(row, 17).value = 0
            ws.cell(row, 25).value = item["m2_caixa"]
            ws.cell(row, 34).value = item["peso_bruto_caixa"]
            for text_column in (2, 8, 13, 26, 27, 28, 29, 30, 31, 32, 40, 51, 52, 53, 57, 58, 59):
                ws.cell(row, text_column).number_format = "@"
        ws.auto_filter.ref = ws.dimensions
        output = OUT_DIR / f"IMPORTACAO_ERP_DELCREDERE_{table}_VINILICO.xlsx"
        book.save(output)

        check = load_workbook(output, read_only=True, data_only=False).active
        rows = list(check.iter_rows(min_row=2, values_only=True))
        mandatory = (1, 3, 5, 6, 7, 12, 13, 14, 15, 18, 19, 23, 24, 25, 26, 27, 28, 29, 30, 31, 38, 39, 50, 56, 57, 58)
        if (check.max_column != 60 or len(rows) != 16 or len({r[1] for r in rows}) != 16 or
                any(not r[index] for r in rows for index in mandatory) or
                any(r[16] != 0 or r[14] <= r[15] for r in rows) or
                any(cell.data_type == "f" for r in check.iter_rows() for cell in r)):
            raise ValueError(f"Falha na validação do arquivo {table}")
        generated[output.name] = sha256(output)

    report = {
        "status": "APROVADO_PARA_TESTE_DE_IMPORTACAO",
        "fonte_pdf": "/home/fernando/Projetos/Work/ARCHOME/ConverterFaltantes/100 TABELA DEL CREDERE SP (4).pdf",
        "fonte_estruturada": str(SOURCE), "fonte_sha256": sha256(SOURCE),
        "vinilicos_por_faixa": 16, "faixas": list(TABLES), "colunas": 60,
        "regra_preco": "PRECO_DA_FAIXA_DEL_CREDERE * 1.0536; DESCONTO_ERP_ZERO",
        "marca_por_faixa": "VILLAGRES 5/10/15/20/25/30",
        "perfil_fiscal": "VINILICO_VALIDADO_PELO_CONTADOR", "peso_bruto": "PESO_BRUTO_CAIXA",
        "arquivos": generated,
        "rodape_bloqueado": "13 itens RP no PDF, mas sem NCM, unidade e tributação homologados; não foram gerados como importação ERP.",
    }
    (OUT_DIR / "RELATORIO_AUDITORIA_DELCREDERE_VINILICO.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT_DIR)


if __name__ == "__main__":
    main()
