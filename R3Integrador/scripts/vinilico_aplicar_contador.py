#!/usr/bin/env python3
"""Aplica o perfil fiscal do contador à importação atual de vinílicos."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


COLUNAS_FISCAIS = [
    13, 18, 19, 20, 21, 26, 27, 28, 29, 30, 31, 32,
    36, 37, 38, 39, 42, 43, 44, 51, 52, 53, 57, 58, 59, 60,
]

CAMPOS_OBRIGATORIOS = {
    2: "CÓDIGO FÁBRICA",
    4: "DESCRIÇÃO COMPLETA",
    6: "GRUPO",
    8: "MARCA",
    13: "NCM",
    14: "UF ORIGEM",
    15: "PREÇO VENDA",
    16: "PREÇO DE FÁBRICA",
    19: "ALIQICMSORIGEM",
    20: "ALIQICMSINTERNA",
    24: "UNIDADE",
    26: "CST",
    27: "ALIQUOTA COFINS CST",
    28: "ALIQUOTA IPI CST",
    29: "ALIQUOTA PIS CST",
    30: "CSOSN",
    31: "CFOP DENTRO",
    32: "CFOP FORA",
    39: "PERCENTUAL ST",
    40: "UNID FABRIL",
    51: "ENQUADRAMENTO IPI",
    57: "ALIQUOTA IBS",
    58: "ALIQUOTA CBS",
    59: "CLASSIFICACAO TRIBUTARIA",
}


def texto(valor: object) -> str:
    return "" if valor is None else str(valor).strip()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--contador", required=True, type=Path)
    parser.add_argument("--atual", required=True, type=Path)
    parser.add_argument("--saida", required=True, type=Path)
    args = parser.parse_args()

    wb_contador = load_workbook(args.contador, data_only=True)
    ws_contador = wb_contador.active
    perfis = {
        tuple(texto(ws_contador.cell(linha, coluna).value) for coluna in COLUNAS_FISCAIS)
        for linha in range(2, ws_contador.max_row + 1)
        if texto(ws_contador.cell(linha, 2).value)
    }
    if len(perfis) != 1:
        raise ValueError(f"Esperado um perfil fiscal único; encontrados: {len(perfis)}")
    perfil = next(iter(perfis))

    wb = load_workbook(args.atual)
    ws = wb.active
    referencias = []
    for linha in range(2, ws.max_row + 1):
        referencia = texto(ws.cell(linha, 2).value)
        if not referencia:
            continue
        referencias.append(referencia)
        for coluna, valor in zip(COLUNAS_FISCAIS, perfil):
            ws.cell(linha, coluna).value = valor

    faltas: list[dict[str, object]] = []
    for linha in range(2, ws.max_row + 1):
        if not texto(ws.cell(linha, 2).value):
            continue
        for coluna, campo in CAMPOS_OBRIGATORIOS.items():
            if not texto(ws.cell(linha, coluna).value):
                faltas.append({"linha": linha, "coluna": coluna, "campo": campo})
        if texto(ws.cell(linha, 27).value) == "0" or texto(ws.cell(linha, 29).value) == "0":
            raise ValueError(f"CST de PIS/COFINS inválido na linha {linha}.")
        if ws.cell(linha, 15).value != ws.cell(linha, 16).value:
            raise ValueError(f"Preço de venda diferente do preço de fábrica na linha {linha}.")
    if faltas:
        raise ValueError(f"Campos obrigatórios ausentes: {faltas}")

    args.saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.saida)

    refs_contador = {
        texto(ws_contador.cell(linha, 2).value)
        for linha in range(2, ws_contador.max_row + 1)
        if texto(ws_contador.cell(linha, 2).value)
    }
    relatorio = {
        "fonte_contador": str(args.contador),
        "fonte_comercial_atual": str(args.atual),
        "quantidade": len(referencias),
        "campos_obrigatorios_ausentes": 0,
        "perfil_fiscal_unico": True,
        "cst_pis": perfil[COLUNAS_FISCAIS.index(29)],
        "cst_cofins": perfil[COLUNAS_FISCAIS.index(27)],
        "referencias_somente_na_versao_contador": sorted(refs_contador - set(referencias)),
        "referencias_somente_na_tabela_atual": sorted(set(referencias) - refs_contador),
        "correspondencia_de_versao": {
            f"LVT9635000{i}": f"LVT9435000{i}" for i in range(1, 5)
        },
        "arquivo": args.saida.name,
        "sha256": hashlib.sha256(args.saida.read_bytes()).hexdigest(),
    }
    auditoria = args.saida.with_name("RELATORIO_AUDITORIA_VINILICO_CONTADOR.json")
    auditoria.write_text(json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")
    print(args.saida.resolve())


if __name__ == "__main__":
    main()
