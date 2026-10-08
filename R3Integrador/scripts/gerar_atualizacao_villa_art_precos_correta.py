#!/usr/bin/env python3
"""Gera atualização de preço Villa Art: preço sugerido na ponta e base de fábrica."""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/home/fernando/Projetos/Work/ARCHOME/Planilhas/VILLAGRES/002 TABELA VILLA ART SP 01.05.26 - ARC HOME.xlsx")
PREVIOUS_IMPORT = ROOT / "Saida" / "IMPORTACAO_ERP_VILLA_ART_20260909.xlsx"
OUTPUT = ROOT / "Saida" / "ATUALIZACAO_VILLA_ART_PRECOS_CORRETOS_20261007.xlsx"
REPORT = ROOT / "Saida" / "RELATORIO_ATUALIZACAO_VILLA_ART_PRECOS_CORRETOS_20261007.json"


def ref(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def rounded(value: object) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    source = load_workbook(SOURCE, read_only=True, data_only=True).active
    imported = load_workbook(PREVIOUS_IMPORT, read_only=True, data_only=True).active
    if str(source.cell(2, 19).value).strip().upper() != "DESCONTO":
        raise ValueError("Coluna S da fonte não identificada como PREÇO DESCONTO")
    if str(source.cell(2, 20).value).strip().upper() != "NA PONTA":
        raise ValueError("Coluna T da fonte não identificada como PREÇO SUGERIDO NA PONTA")

    erp_rows = {ref(imported.cell(row, 2).value): row for row in range(2, imported.max_row + 1)}
    source_rows: dict[str, tuple[int, Decimal, Decimal]] = {}
    for row in range(3, source.max_row + 1):
        code = ref(source.cell(row, 2).value)
        if not code:
            continue
        if code in source_rows:
            raise ValueError(f"Referência duplicada na fonte: {code}")
        factory = rounded(source.cell(row, 19).value)
        sale = rounded(source.cell(row, 20).value)
        source_rows[code] = (row, factory, sale)

    if len(source_rows) != 25:
        raise ValueError(f"Esperadas 25 referências na fonte, encontradas {len(source_rows)}")
    if set(source_rows) != set(erp_rows):
        raise ValueError("As referências da tabela não coincidem com os 25 itens importados no ERP")

    workbook = Workbook()
    ws = workbook.active
    ws.title = "ATUALIZACAO"
    headers = ["MARCA", "REFERENCIA", "PREÇO VENDA", "PREÇO DE FÁBRICA", "DESCRICAO", "UNID FABRIL", "MODELO", "COR"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    items = []
    for code, (source_row, factory, sale) in source_rows.items():
        erp_row = erp_rows[code]
        ws.append([
            imported.cell(erp_row, 8).value,
            code,
            float(sale),
            float(factory),
            imported.cell(erp_row, 4).value,
            imported.cell(erp_row, 40).value,
            imported.cell(erp_row, 10).value,
            imported.cell(erp_row, 12).value,
        ])
        previous_sale = rounded(imported.cell(erp_row, 15).value)
        previous_factory = rounded(imported.cell(erp_row, 16).value)
        items.append({
            "referencia": code,
            "preco_tabela": float(rounded(source.cell(source_row, 18).value)),
            "preco_desconto_usado_como_fabrica": float(factory),
            "preco_sugerido_na_ponta_usado_como_venda": float(sale),
            "preco_venda_arquivo_anterior": float(previous_sale),
            "preco_fabrica_arquivo_anterior": float(previous_factory),
        })
    for cell in ws["C"][1:] + ws["D"][1:]:
        cell.number_format = 'R$ #,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in {"A": 16, "B": 18, "C": 18, "D": 20, "E": 62, "F": 14, "G": 18, "H": 18}.items():
        ws.column_dimensions[col].width = width
    workbook.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=True, data_only=True).active
    rows = list(check.iter_rows(min_row=2, values_only=True))
    if (check.max_column != 8 or len(rows) != 25 or
            len({row[1] for row in rows}) != 25 or
            any(not row[2] or not row[3] for row in rows)):
        raise ValueError("Falha na validação da planilha de atualização")
    ref_row = next(row for row in rows if row[1] == "123049")
    if rounded(ref_row[2]) != Decimal("155.68") or rounded(ref_row[3]) != Decimal("91.31"):
        raise ValueError(f"Preço de validação incorreto para 123049: {ref_row[2]} / {ref_row[3]}")

    report = {
        "status": "ATUALIZACAO_PREPARADA_PARA_REVISAO_E_IMPORTACAO_ERP",
        "fonte": str(SOURCE),
        "fonte_sha256": sha256(SOURCE),
        "arquivo_atualizacao": str(OUTPUT),
        "arquivo_sha256": sha256(OUTPUT),
        "quantidade": len(rows),
        "layout": headers,
        "regra": "PREÇO DE FÁBRICA = PREÇO DESCONTO (coluna S, arredondado a 2 casas); PREÇO VENDA = PREÇO SUGERIDO NA PONTA (coluna T, arredondado a 2 casas). Sem adicionar IPI/ST/markup ao preço de venda.",
        "erp": "Preço de fábrica fica como base para o ERP calcular os tributos já cadastrados; atualização não altera campos fiscais.",
        "validacao_123049": {"preco_tabela": 115.58, "preco_desconto_fabrica": 91.31, "preco_sugerido_venda": 155.68},
        "itens": items,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT)
    print(REPORT)


if __name__ == "__main__":
    main()
