#!/usr/bin/env python3
"""Audita as saídas geradas para as correções solicitadas por Mônica."""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def integrity(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        return archive.testzip() is None


def values(worksheet):
    # Algumas planilhas têm formatação aplicada até a última linha do Excel.
    # As fontes deste fluxo têm no máximo 1.141 produtos; limitar a leitura
    # evita varrer mais de um milhão de linhas vazias.
    return [
        row
        for row in worksheet.iter_rows(min_row=2, max_row=min(worksheet.max_row, 10_000), values_only=True)
        if any(v not in (None, "") for v in row)
    ]


def formula_count(worksheet, last_row: int) -> int:
    return sum(
        1
        for row in worksheet.iter_rows(min_row=1, max_row=last_row, max_col=worksheet.max_column)
        for cell in row
        if cell.data_type == "f"
    )


def audit_adama(path: Path) -> dict[str, object]:
    worksheet = load_workbook(path, read_only=True, data_only=False)["IMPORTACAO_ERP"]
    rows = values(worksheet)
    row2 = rows[0]
    return {
        "arquivo": path.name,
        "sha256": sha256(path),
        "xlsx_integro": integrity(path),
        "linhas": len(rows),
        "colunas": worksheet.max_column,
        "formulas": formula_count(worksheet, len(rows) + 1),
        "referencia_linha_2": row2[1],
        "ncm_linha_2": row2[12],
        "icms_origem_linha_2": row2[18],
        "icms_interna_linha_2": row2[19],
        "icms_linha_2_preenchido": row2[18] == 18 and row2[19] == 18,
    }


def audit_studio(path: Path) -> dict[str, object]:
    worksheet = load_workbook(path, read_only=True, data_only=False)["IMPORTACAO_ERP"]
    rows = values(worksheet)
    csosn = {str(row[29]) for row in rows}
    formats = {worksheet.cell(row, 30).number_format for row in range(2, len(rows) + 2)}
    return {
        "arquivo": path.name,
        "sha256": sha256(path),
        "xlsx_integro": integrity(path),
        "linhas": len(rows),
        "colunas": worksheet.max_column,
        "formulas": formula_count(worksheet, len(rows) + 1),
        "csosn_distintos": sorted(csosn),
        "csosn_todos_500": csosn == {"500"},
        "formato_csosn_texto": formats == {"@"},
    }


def normalized_code(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def audit_nina(path: Path, source: Path) -> dict[str, object]:
    # A auditoria consulta o formato de diversas células; modo normal evita
    # revarrer o XML inteiro para cada consulta no modo somente leitura.
    workbook = load_workbook(path, read_only=False, data_only=False)
    worksheet = workbook["IMPORTACAO_ERP"]
    rows = values(worksheet)
    source_ws = load_workbook(source, read_only=True, data_only=True)["Coleção Completa"]
    source_rows = [
        row for row in source_ws.iter_rows(min_row=2, values_only=True)
        if normalized_code(row[2]) and normalized_code(row[3])
    ]
    price_errors = []
    source_code_errors = []
    for index, (output, origin) in enumerate(zip(rows, source_rows), start=2):
        expected_code = "" if normalized_code(origin[2]).casefold() == "depende do raio" else normalized_code(origin[2])
        if normalized_code(output[1]) != expected_code:
            source_code_errors.append((index, normalized_code(output[1]), expected_code))
        if abs(Decimal(str(output[14])) - Decimal(str(origin[12]))) > Decimal("0.000001"):
            price_errors.append((index, output[1], output[14], origin[12]))

    fiscal = {
        "ipi": {row[17] for row in rows},
        "icms_origem": {row[18] for row in rows},
        "icms_interna": {row[19] for row in rows},
        "iva": {row[20] for row in rows},
        "embalagem_venda": {row[24] for row in rows},
        "cst": {str(row[25]) for row in rows},
        "cofins_cst": {str(row[26]) for row in rows},
        "ipi_cst": {str(row[27]) for row in rows},
        "pis_cst": {str(row[28]) for row in rows},
        "csosn": {str(row[29]) for row in rows},
        "cfop_dentro": {str(row[30]) for row in rows},
        "cfop_fora": {str(row[31]) for row in rows},
        "enquadramento_ipi": {str(row[50]) for row in rows},
        "pis_origem": {str(row[51]) for row in rows},
        "cofins_origem": {str(row[52]) for row in rows},
        "ibs": {str(row[56]) for row in rows},
        "cbs": {str(row[57]) for row in rows},
        "classificacao": {str(row[58]) for row in rows},
    }
    required = [25, 26, 27, 28, 29, 30, 31, 50, 51, 52, 56, 57, 58]
    missing = {str(column + 1): sum(row[column] in (None, "") for row in rows) for column in required}
    tax_formats = {
        column: {worksheet.cell(row, column).number_format for row in range(2, len(rows) + 2)}
        for column in (13, 26, 27, 28, 29, 30, 31, 32, 51, 59)
    }
    nonempty_refs = [normalized_code(row[1]) for row in rows if normalized_code(row[1])]
    return {
        "arquivo": path.name,
        "sha256": sha256(path),
        "xlsx_integro": integrity(path),
        "linhas": len(rows),
        "colunas": worksheet.max_column,
        "formulas": formula_count(worksheet, len(rows) + 1),
        "linhas_fonte_rev04": len(source_rows),
        "precos_divergentes_rev04": len(price_errors),
        "codigos_divergentes_rev04": len(source_code_errors),
        "referencias_preenchidas_unicas": len(nonempty_refs) == len(set(nonempty_refs)),
        "codigos_fabrica_vazios": sum(not normalized_code(row[1]) for row in rows),
        "campos_obrigatorios_ausentes": missing,
        "valores_fiscais_distintos": {key: sorted(value) for key, value in fiscal.items()},
        "codigos_formatados_como_texto": {
            str(column): formats == {"@"} for column, formats in tax_formats.items()
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--saida", type=Path, required=True)
    parser.add_argument("--nina-rev04", type=Path, required=True)
    args = parser.parse_args()
    adama = args.saida / "IMPORTACAO_ERP_ADAMA_CORRIGIDA_MONICA_20260907.xlsx"
    studio = args.saida / "IMPORTACAO_ERP_STUDIO_MORANDIN_CORRIGIDA_MONICA_20260907.xlsx"
    nina = Path("Saida/IMPORTACAO_ERP_NINA_MARTINELLI_IMPOSTOS_CONTADOR_20260907_123026.xlsx")
    report = {
        "adama": audit_adama(adama),
        "studio_morandin": audit_studio(studio),
        "nina_martinelli": audit_nina(nina, args.nina_rev04),
    }
    (args.saida / "RELATORIO_AUDITORIA_COMPLETO_MONICA_20260907.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(args.saida / "RELATORIO_AUDITORIA_COMPLETO_MONICA_20260907.json")


if __name__ == "__main__":
    main()
