#!/usr/bin/env python3
"""Audita, corrige, expande e importa os dados VillaCol.

O pipeline usa as planilhas DEL5 a DEL30 aprovadas pelo contador como fonte
dos preços de fábrica e da parametrização fiscal. O PDF de revenda é a fonte
das referências, descrições, embalagens, quantidades e preços de revenda.

Nenhum arquivo de origem é sobrescrito. As planilhas corrigidas, a importação
de revenda e os relatórios de auditoria são gravados em uma pasta de saída.
Antes de alterar o SQLite, o banco é copiado para essa mesma pasta.
"""

from __future__ import annotations

import argparse
import json
import logging
import math
import os
import re
import shutil
import sqlite3
import sys
from collections import Counter
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


LOGGER = logging.getLogger("villacol")

SHEET_NAME = "IMPORTACAO_ERP"
EXPECTED_DEL_LEVELS = (5, 10, 15, 20, 25, 30)
EXPECTED_BASE_REFERENCES = 41
EXPECTED_SALES_SKUS = 63
EXPECTED_DEL_ITEMS = 10

DEFAULT_INPUT_DIR = Path("/home/fernando/Projetos/Work/ARCHOME/Planilhas/VillaCol")
DEFAULT_PDF_NAME = "TABELA VAREJO 04-05-2026.pdf"

PRODUCT_COLUMNS = {
    "CÓDIGO INTERNO",
    "CÓDIGO FÁBRICA",
    "CODIGOBARRAS",
    "DESCRIÇÃO COMPLETA",
    "DESCRIÇÃO COMERCIAL",
    "GRUPO",
    "SUBGRUPO",
    "MARCA",
    "LINHA",
    "MODELO",
    "VOLTAGEM",
    "COR",
    "PREÇO VENDA",
    "PREÇO DE FÁBRICA",
    "UNIDADE",
    "QTDE EMBALAGEM DE VENDA",
    "PESOLIQ",
    "PESOBRUTO",
    "QTDE EMBALAGEM DE COMPRA",
    "UNID FABRIL",
    "OBSERVACAO",
}

CODE_WIDTHS = {
    "NCM": 8,
    "CST": 3,
    "ALIQUOTA COFINS CST": 2,
    "ALIQUOTA IPI CST": 2,
    "ALIQUOTA PIS CST": 2,
    "CSOSN": 3,
    "CFOP DENTRO": 4,
    "CFOP FORA": 4,
    "ENQUADRAMENTO IPI": 3,
    "CLASSIFICACAO TRIBUTARIA": 6,
}

# A ordem corresponde exatamente às 60 colunas do layout ERP.
DB_COLUMNS = (
    "CodigoInterno", "CodigoFabrica", "CodigoBarras", "DescricaoCompleta",
    "DescricaoComercial", "Grupo", "SubGrupo", "Marca", "Linha", "Modelo",
    "Voltagem", "Cor", "Ncm", "UfOrigem", "PrecoVenda", "PrecoFabrica",
    "DescontoPercentual", "IpiPercentual", "AliqIcmsOrigem", "AliqIcmsInterna",
    "Iva", "FreteReais", "FretePercentual", "Unidade", "QtdeEmbalagemVenda",
    "Cst", "AliquotaCofinsCst", "AliquotaIpiCst", "AliquotaPisCst", "Csosn",
    "CfopDentro", "CfopFora", "PesoLiquido", "PesoBruto",
    "QtdeEmbalagemCompra", "ValorPi", "AliquotaCofins", "AliquotaPis",
    "PercentualSt", "UnidFabril", "Observacao", "DiferencaIcms",
    "ReducaoBaseIcms", "ReducaoBaseSt", "RetencaoPis", "RetencaoCofins",
    "RetencaoCsll", "RetencaoIrrf", "RetencaoPrevSocial", "Localizacao",
    "EnquadramentoIpi", "AliquotaPisOrigem", "AliquotaCofinsOrigem", "Imagem",
    "EstoqueMinimo", "EstoqueMaximo", "AliquotaIbs", "AliquotaCbs",
    "ClassificacaoTributaria", "CodigoBeneficio",
)


@dataclass(frozen=True)
class SourceProduct:
    """Um SKU comercial extraído do PDF."""

    sku: str
    pdf_reference: str
    derivation: str
    description: str
    category: str
    subgroup: str
    color: str
    weight_kg: Decimal
    unit: str
    package_quantity: int
    pallet_quantity: int
    resale_price: Decimal
    price_class: str


@dataclass(frozen=True)
class AuditResult:
    source_rows_per_del: dict[int, int]
    missing_factory_codes_per_del: dict[int, int]
    del30_max_implied_base_spread: float
    del30_prices_consistent: bool
    normalized_tax_fields: dict[str, Any]
    pdf_base_references: int
    pdf_sales_skus: int


def clean_text(value: Any) -> str:
    """Converte valores vazios/NaN em string vazia e normaliza espaços."""

    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    text = str(value).replace("\u00a0", " ").strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return re.sub(r"\s+", " ", text)


def decimal_value(value: Any) -> Decimal:
    """Interpreta valores monetários brasileiros e números do Excel."""

    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and math.isnan(value):
            return Decimal("0")
        return Decimal(str(value))

    text = clean_text(value).replace("R$", "").replace(" ", "")
    if not text:
        return Decimal("0")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Valor decimal inválido: {value!r}") from exc


def money(value: Any) -> Decimal:
    return decimal_value(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def first_integer(value: Any, default: int = 0) -> int:
    match = re.search(r"\d+", clean_text(value))
    return int(match.group()) if match else default


def normalize_code(value: Any, width: int) -> str:
    text = clean_text(value)
    if not text:
        return ""
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        text = str(int(Decimal(text)))
    return text.zfill(width)


def most_common_clean(values: Iterable[Any]) -> Any:
    cleaned = [clean_text(value) for value in values]
    counter = Counter(cleaned)
    if not counter:
        return ""
    # Em empate, prefere o valor não vazio e depois a primeira ocorrência.
    best_count = max(counter.values())
    candidates = {value for value, count in counter.items() if count == best_count}
    for value in cleaned:
        if value in candidates and value:
            return value
    return cleaned[0]


class VillaColPdfExtractor:
    """Extrai os produtos da tabela de revenda preservando a ordem do PDF."""

    ARG_DESCRIPTIONS = {
        "3X1CZA": ("ARGAMASSA 3 EM 1 - CINZA", "3 EM 1"),
        "7X1BRA": ("ARGAMASSA 7 EM 1 - BRANCA", "7 EM 1"),
        "7X1CZA": ("ARGAMASSA 7 EM 1 - CINZA", "7 EM 1"),
        "GRD": ("ARGAMASSA ACIII - GRD", "ACIII"),
    }

    def extract(self, pdf_path: Path) -> list[SourceProduct]:
        if not pdf_path.exists():
            raise FileNotFoundError(pdf_path)

        reader = PdfReader(str(pdf_path))
        if len(reader.pages) != 1:
            raise ValueError(f"Esperava PDF de 1 página; encontrado: {len(reader.pages)}")

        with pdfplumber.open(pdf_path) as pdf:
            page = pdf.pages[0]
            tables = page.extract_tables()
            if len(tables) != 1:
                raise ValueError(f"Esperava 1 tabela no PDF; encontrado: {len(tables)}")
            table = tables[0]
            word_prices = self._extract_word_prices(page)

        products: list[SourceProduct] = []
        res_rows = [row for row in table if self._reference(row).startswith("RES")]
        fix_rows = [row for row in table if self._reference(row).startswith("FIX")]

        res_price = self._first_nonzero(row[6] for row in res_rows)
        res_pallet = max((first_integer(row[5]) for row in res_rows), default=0)
        fix_price = self._first_nonzero(row[6] for row in fix_rows)
        fix_pallet = max((first_integer(row[5]) for row in fix_rows), default=0)

        for row in table:
            reference = self._reference(row)
            if not reference:
                continue

            if reference in self.ARG_DESCRIPTIONS:
                description, subgroup = self.ARG_DESCRIPTIONS[reference]
                color = clean_text(row[1]).replace(" ", "")
                products.append(SourceProduct(
                    sku=reference,
                    pdf_reference=reference,
                    derivation=clean_text(row[4]),
                    description=description,
                    category="ARGAMASSAS",
                    subgroup=subgroup,
                    color=color,
                    weight_kg=Decimal("20"),
                    unit="SC",
                    package_quantity=1,
                    pallet_quantity=first_integer(row[5]),
                    resale_price=money(row[6]),
                    price_class=description,
                ))
                continue

            description = clean_text(row[2]).upper()
            if reference.startswith("RES"):
                products.append(SourceProduct(
                    sku=reference,
                    pdf_reference=reference,
                    derivation="REJ1KG",
                    description=description,
                    category="REJUNTES",
                    subgroup="RESINADO",
                    color=self._color(description, "REJUNTE RESINADO"),
                    weight_kg=Decimal("1"),
                    unit="SC",
                    package_quantity=15,
                    pallet_quantity=res_pallet,
                    resale_price=res_price,
                    price_class="REJUNTE RESINADO 1KG",
                ))
                continue

            if reference.startswith("FIX"):
                products.append(SourceProduct(
                    sku=reference,
                    pdf_reference=reference,
                    derivation="FIX15KG",
                    description=description,
                    category="REJUNTES",
                    subgroup="FIX",
                    color=self._color(description, "FIX"),
                    weight_kg=Decimal("15"),
                    unit="SC",
                    package_quantity=1,
                    pallet_quantity=fix_pallet,
                    resale_price=fix_price,
                    price_class="FIX ASSENTA E REJUNTA 15KG",
                ))
                continue

            if reference.startswith("ACR"):
                row_prices = [money(row[6]), money(row[7])]
                if not all(row_prices):
                    row_prices = word_prices.get(reference, row_prices)
                if len(row_prices) != 2 or not all(row_prices):
                    raise ValueError(f"Dois preços não encontrados para {reference}: {row_prices}")

                color = self._color(description, "REJUNTE ACRÍLICO")
                special = row_prices[0] == Decimal("45.90")
                for weight, derivation, pallet, resale_price in (
                    (1, "PTA1KG", 270, row_prices[0]),
                    (2, "PTA2KG", 180, row_prices[1]),
                ):
                    price_class = (
                        f"REJUNTE ACRILICO (CORES ESPECIAIS) {weight}KG"
                        if special
                        else f"REJUNTE ACRILICO {weight}KG"
                    )
                    products.append(SourceProduct(
                        sku=f"{reference}-{derivation}",
                        pdf_reference=reference,
                        derivation=derivation,
                        description=f"{description} {weight}KG",
                        category="REJUNTES",
                        subgroup="ACRÍLICO",
                        color=color,
                        weight_kg=Decimal(weight),
                        unit="PT",
                        package_quantity=6,
                        pallet_quantity=pallet,
                        resale_price=resale_price,
                        price_class=price_class,
                    ))

        self._validate(products)
        return products

    @staticmethod
    def _reference(row: Sequence[Any]) -> str:
        if len(row) < 4:
            return ""
        reference = clean_text(row[3]).upper()
        return reference if re.fullmatch(r"(?:3X1CZA|7X1BRA|7X1CZA|GRD|RES\d{3}|FIX\d{3}|ACR\d{3})", reference) else ""

    @staticmethod
    def _color(description: str, prefix: str) -> str:
        return clean_text(description.removeprefix(prefix)).upper()

    @staticmethod
    def _first_nonzero(values: Iterable[Any]) -> Decimal:
        for value in values:
            parsed = money(value)
            if parsed:
                return parsed
        return Decimal("0")

    @staticmethod
    def _extract_word_prices(page: Any) -> dict[str, list[Decimal]]:
        words = page.extract_words(x_tolerance=2, y_tolerance=2, use_text_flow=False)
        references = {
            word["text"].upper(): float(word["top"])
            for word in words
            if re.fullmatch(r"ACR\d{3}", word["text"].upper())
        }
        prices: dict[str, list[Decimal]] = {}
        for reference, top in references.items():
            values = [
                money(word["text"])
                for word in words
                if abs(float(word["top"]) - top) <= 1.1
                and float(word["x0"]) >= 460
                and re.fullmatch(r"\d+,\d{2}", word["text"])
            ]
            if values:
                prices[reference] = values
        return prices

    @staticmethod
    def _validate(products: Sequence[SourceProduct]) -> None:
        base_references = {product.pdf_reference for product in products}
        skus = [product.sku for product in products]
        if len(base_references) != EXPECTED_BASE_REFERENCES:
            raise ValueError(
                f"Referências do PDF: {len(base_references)}; esperado: {EXPECTED_BASE_REFERENCES}"
            )
        if len(skus) != EXPECTED_SALES_SKUS:
            raise ValueError(f"SKUs extraídos: {len(skus)}; esperado: {EXPECTED_SALES_SKUS}")
        duplicates = [sku for sku, count in Counter(skus).items() if count > 1]
        if duplicates:
            raise ValueError(f"SKUs duplicados: {duplicates}")
        if any(product.resale_price <= 0 for product in products):
            raise ValueError("Há produto sem preço de revenda positivo no PDF.")


class ApprovedSheets:
    """Lê as seis planilhas e fornece preços/fiscalidade consolidados."""

    # As tabelas Del Credere são deliberadamente mantidas com os dez itens
    # agregados históricos. Somente a tabela de revenda é expandida com todos
    # os SKUs do PDF.
    DEL_ITEM_RULES: dict[str, dict[str, Any]] = {
        "ARGAMASSA 3 EM 1 - CINZA": {
            "code": "3X1CZA", "group": "ARGAMASSAS", "subgroup": "3 EM 1",
            "model": "20KG", "color": "CINZA", "unit": "SC", "package": 1, "weight": 20,
        },
        "ARGAMASSA 7 EM 1 - BRANCA": {
            "code": "7X1BRA", "group": "ARGAMASSAS", "subgroup": "7 EM 1",
            "model": "20KG", "color": "BRANCA", "unit": "SC", "package": 1, "weight": 20,
        },
        "ARGAMASSA 7 EM 1 - CINZA": {
            "code": "7X1CZA", "group": "ARGAMASSAS", "subgroup": "7 EM 1",
            "model": "20KG", "color": "CINZA", "unit": "SC", "package": 1, "weight": 20,
        },
        "ARGAMASSA ACIII - GRD": {
            "code": "GRD", "group": "ARGAMASSAS", "subgroup": "ACIII",
            "model": "20KG", "color": "CINZA", "unit": "SC", "package": 1, "weight": 20,
        },
        "REJUNTE ACRILICO 2KG": {
            "code": "PTA2KG", "group": "REJUNTES", "subgroup": "ACRÍLICO",
            "model": "2KG", "color": "CORES PADRÃO", "unit": "PT", "package": 6, "weight": 2,
        },
        "REJUNTE ACRILICO (CORES ESPECIAIS) 2KG": {
            "code": "PTA2KG-ESPECIAL", "group": "REJUNTES", "subgroup": "ACRÍLICO",
            "model": "2KG", "color": "CORES ESPECIAIS", "unit": "PT", "package": 6, "weight": 2,
        },
        "REJUNTE ACRILICO 1KG": {
            "code": "PTA1KG", "group": "REJUNTES", "subgroup": "ACRÍLICO",
            "model": "1KG", "color": "CORES PADRÃO", "unit": "PT", "package": 6, "weight": 1,
        },
        "REJUNTE ACRILICO (CORES ESPECIAIS) 1KG": {
            "code": "PTA1KG-ESPECIAL", "group": "REJUNTES", "subgroup": "ACRÍLICO",
            "model": "1KG", "color": "CORES ESPECIAIS", "unit": "PT", "package": 6, "weight": 1,
        },
        "REJUNTE RESINADO 1KG": {
            "code": "REJ1KG", "group": "REJUNTES", "subgroup": "RESINADO",
            "model": "1KG", "color": "CORES", "unit": "SC", "package": 15, "weight": 1,
        },
        "FIX ASSENTA E REJUNTA 15KG": {
            "code": "FIX15KG", "group": "REJUNTES", "subgroup": "FIX",
            "model": "15KG", "color": "CORES", "unit": "SC", "package": 1, "weight": 15,
        },
    }

    def __init__(self, input_dir: Path):
        self.paths = self._discover(input_dir)
        self.frames = {
            level: pd.read_excel(path, sheet_name=SHEET_NAME, dtype=object, keep_default_na=False)
            for level, path in self.paths.items()
        }
        self.headers = list(self.frames[EXPECTED_DEL_LEVELS[0]].columns)
        self._validate_layouts()
        self.canonical_defaults = self._canonical_defaults()
        self.factory_prices = self._factory_prices()

    @staticmethod
    def _discover(input_dir: Path) -> dict[int, Path]:
        paths: dict[int, Path] = {}
        for path in input_dir.glob("*.xlsx"):
            match = re.search(r"DELCREDERE_DEL(5|10|15|20|25|30)_VILLACOL", path.name, re.I)
            if match:
                paths[int(match.group(1))] = path
        missing = sorted(set(EXPECTED_DEL_LEVELS) - set(paths))
        if missing:
            raise FileNotFoundError(f"Planilhas DEL ausentes: {missing}")
        return dict(sorted(paths.items()))

    def _validate_layouts(self) -> None:
        if len(self.headers) != 60:
            raise ValueError(f"Layout ERP deve ter 60 colunas; encontrado: {len(self.headers)}")
        for level, frame in self.frames.items():
            if list(frame.columns) != self.headers:
                raise ValueError(f"Cabeçalho da DEL{level} difere das demais planilhas.")
            if frame.empty:
                raise ValueError(f"Planilha DEL{level} vazia.")

    def _canonical_defaults(self) -> dict[str, Any]:
        defaults: dict[str, Any] = {header: "" for header in self.headers}
        for header in self.headers:
            if header in PRODUCT_COLUMNS:
                continue
            values = [value for frame in self.frames.values() for value in frame[header].tolist()]
            defaults[header] = most_common_clean(values)

        for column, width in CODE_WIDTHS.items():
            defaults[column] = normalize_code(defaults[column], width)
        return defaults

    def _factory_prices(self) -> dict[int, dict[str, Decimal]]:
        result: dict[int, dict[str, Decimal]] = {}
        for level, frame in self.frames.items():
            lookup = {
                clean_text(row["DESCRIÇÃO COMPLETA"]).upper(): money(row["PREÇO DE FÁBRICA"])
                for _, row in frame.iterrows()
            }
            if len(lookup) != 10 or any(not key for key in lookup):
                raise ValueError(f"DEL{level} não contém as 10 classes de preço esperadas.")
            result[level] = lookup
        return result

    def audit(self, products: Sequence[SourceProduct]) -> AuditResult:
        implied: dict[str, list[Decimal]] = {}
        for level, prices in self.factory_prices.items():
            for description, price in prices.items():
                implied.setdefault(description, []).append(price * (Decimal("1") - Decimal(level) / 100))
        spreads = [max(values) - min(values) for values in implied.values()]
        max_spread = max(spreads, default=Decimal("0"))

        return AuditResult(
            source_rows_per_del={level: len(frame) for level, frame in self.frames.items()},
            missing_factory_codes_per_del={
                level: int(sum(not clean_text(value) for value in frame["CÓDIGO FÁBRICA"]))
                for level, frame in self.frames.items()
            },
            del30_max_implied_base_spread=float(max_spread),
            del30_prices_consistent=max_spread <= Decimal("0.01"),
            normalized_tax_fields={
                key: self.canonical_defaults[key]
                for key in (
                    "NCM", "UF ORIGEM", "CST", "ALIQUOTA COFINS CST",
                    "ALIQUOTA IPI CST", "ALIQUOTA PIS CST", "CSOSN",
                    "CFOP DENTRO", "CFOP FORA", "ENQUADRAMENTO IPI",
                    "ALIQUOTA PIS ORIGEM", "ALIQUOTA COFINS ORIGEM",
                    "ALIQUOTA IBS", "ALIQUOTA CBS", "CLASSIFICACAO TRIBUTARIA",
                )
            },
            pdf_base_references=len({product.pdf_reference for product in products}),
            pdf_sales_skus=len(products),
        )

    def build_del_frame(self, level: int) -> pd.DataFrame:
        """Corrige os dez itens históricos sem expandi-los por cor/referência."""

        rows: list[dict[str, Any]] = []
        for _, original in self.frames[level].iterrows():
            description = clean_text(original["DESCRIÇÃO COMPLETA"])
            rule = self.DEL_ITEM_RULES.get(description.upper())
            if rule is None:
                raise ValueError(f"Regra não encontrada para item histórico: {description}")

            row = dict(self.canonical_defaults)
            row.update({
                "CÓDIGO INTERNO": "",
                "CÓDIGO FÁBRICA": rule["code"],
                "CODIGOBARRAS": "",
                "DESCRIÇÃO COMPLETA": description,
                "DESCRIÇÃO COMERCIAL": clean_text(original["DESCRIÇÃO COMERCIAL"]) or description,
                "GRUPO": rule["group"],
                "SUBGRUPO": rule["subgroup"],
                "MARCA": "VILLACOL",
                "LINHA": "DELCREDERE",
                "MODELO": rule["model"],
                "VOLTAGEM": "",
                "COR": rule["color"],
                "PREÇO VENDA": Decimal("0"),
                "PREÇO DE FÁBRICA": money(original["PREÇO DE FÁBRICA"]),
                "UNIDADE": rule["unit"],
                "QTDE EMBALAGEM DE VENDA": Decimal("0"),
                "PESOLIQ": Decimal(rule["weight"]),
                "PESOBRUTO": Decimal("0"),
                "QTDE EMBALAGEM DE COMPRA": rule["package"],
                "UNID FABRIL": rule["unit"],
                "OBSERVACAO": "",
            })
            rows.append(row)

        frame = pd.DataFrame(rows, columns=self.headers)
        self._validate_output(frame, EXPECTED_DEL_ITEMS)
        return frame

    def build_resale_frame(self, products: Sequence[SourceProduct]) -> pd.DataFrame:
        """Expande todos os produtos/embalagens encontrados no PDF de revenda."""

        rows: list[dict[str, Any]] = []
        for product in products:
            row = dict(self.canonical_defaults)
            row.update({
                "CÓDIGO INTERNO": "",
                "CÓDIGO FÁBRICA": product.sku,
                "CODIGOBARRAS": "",
                "DESCRIÇÃO COMPLETA": product.description,
                "DESCRIÇÃO COMERCIAL": product.description,
                "GRUPO": product.category,
                "SUBGRUPO": product.subgroup,
                "MARCA": "VILLACOL",
                "LINHA": "REVENDA",
                "MODELO": f"{product.weight_kg:g}KG",
                "VOLTAGEM": "",
                "COR": product.color,
                "PREÇO VENDA": Decimal("0"),
                "PREÇO DE FÁBRICA": product.resale_price,
                "UNIDADE": product.unit,
                "QTDE EMBALAGEM DE VENDA": Decimal("0"),
                "PESOLIQ": product.weight_kg,
                # O PDF informa peso nominal, não o peso bruto da embalagem.
                "PESOBRUTO": Decimal("0"),
                "QTDE EMBALAGEM DE COMPRA": product.package_quantity,
                "UNID FABRIL": product.unit,
                "OBSERVACAO": "",
            })
            rows.append(row)

        frame = pd.DataFrame(rows, columns=self.headers)
        self._validate_output(frame, EXPECTED_SALES_SKUS)
        return frame

    @staticmethod
    def _validate_output(frame: pd.DataFrame, expected_rows: int) -> None:
        if len(frame) != expected_rows:
            raise ValueError(f"Planilha gerada com {len(frame)} itens; esperado: {expected_rows}")
        if frame["CÓDIGO FÁBRICA"].duplicated().any():
            raise ValueError("A planilha gerada contém códigos de fábrica duplicados.")
        if any(not clean_text(value) for value in frame["CÓDIGO FÁBRICA"]):
            raise ValueError("A planilha gerada contém código de fábrica vazio.")
        if any(decimal_value(value) != 0 for value in frame["PREÇO VENDA"]):
            raise ValueError("PREÇO VENDA deve ser zero em todos os itens.")
        for column in ("GRUPO", "SUBGRUPO", "OBSERVACAO"):
            if any(clean_text(value).lower() == "nan" for value in frame[column]):
                raise ValueError(f"A coluna {column} contém NaN textual.")


class ExcelWriter:
    """Gera Excel preservando cabeçalho e estilo da planilha aprovada."""

    TEXT_COLUMNS = {
        "CÓDIGO INTERNO", "CÓDIGO FÁBRICA", "CODIGOBARRAS", "NCM", "CST",
        "ALIQUOTA COFINS CST", "ALIQUOTA IPI CST", "ALIQUOTA PIS CST", "CSOSN",
        "CFOP DENTRO", "CFOP FORA", "ENQUADRAMENTO IPI",
        "CLASSIFICACAO TRIBUTARIA", "CODIGO BENEFICIO",
    }

    def write(self, template_path: Path, frame: pd.DataFrame, output_path: Path) -> None:
        workbook = load_workbook(template_path)
        worksheet = workbook[SHEET_NAME]
        styles = [copy(worksheet.cell(2, column)._style) for column in range(1, 61)]
        alignments = [copy(worksheet.cell(2, column).alignment) for column in range(1, 61)]
        row_height = worksheet.row_dimensions[2].height

        if worksheet.max_row > 1:
            worksheet.delete_rows(2, worksheet.max_row - 1)

        for row_number, values in enumerate(frame.itertuples(index=False, name=None), start=2):
            worksheet.row_dimensions[row_number].height = row_height
            for column_number, value in enumerate(values, start=1):
                header = frame.columns[column_number - 1]
                cell = worksheet.cell(row_number, column_number)
                cell._style = copy(styles[column_number - 1])
                cell.alignment = copy(alignments[column_number - 1])
                if isinstance(value, Decimal):
                    value = float(value)
                value = None if clean_text(value) == "" else value
                cell.value = value
                if header in self.TEXT_COLUMNS and value is not None:
                    cell.value = str(value)
                    cell.number_format = "@"
                elif header in {"PREÇO VENDA", "PREÇO DE FÁBRICA"}:
                    cell.number_format = '#,##0.00'

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:BH{len(frame) + 1}"
        worksheet.sheet_view.showGridLines = False
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = output_path.with_suffix(".tmp.xlsx")
        workbook.save(temporary)
        os.replace(temporary, output_path)


class DatabaseLoader:
    """Persiste o layout ERP e os metadados extraídos no SQLite."""

    ORIGINS = tuple([f"VILLACOL_DEL{level}" for level in EXPECTED_DEL_LEVELS] + ["VILLACOL_REVENDA"])

    def load(
        self,
        database_path: Path,
        output_dir: Path,
        frames: Mapping[str, pd.DataFrame],
        products: Sequence[SourceProduct],
        source_pdf: Path,
    ) -> dict[str, int]:
        if not database_path.exists():
            raise FileNotFoundError(database_path)
        backup_path = output_dir / f"{database_path.stem}_ANTES_VILLACOL{database_path.suffix}.bak"
        # Não sobrescreve o primeiro backup em reexecuções idempotentes.
        if not backup_path.exists():
            shutil.copy2(database_path, backup_path)

        with sqlite3.connect(database_path) as connection:
            self._validate_product_schema(connection)
            connection.execute("BEGIN IMMEDIATE")
            try:
                placeholders = ",".join("?" for _ in (*DB_COLUMNS, "TabelaOrigem"))
                columns_sql = ",".join(f'"{column}"' for column in (*DB_COLUMNS, "TabelaOrigem"))
                connection.executemany(
                    "DELETE FROM Produtos WHERE TabelaOrigem = ?",
                    [(origin,) for origin in self.ORIGINS],
                )

                for origin, frame in frames.items():
                    records = [
                        tuple(self._sqlite_value(value) for value in row) + (origin,)
                        for row in frame.itertuples(index=False, name=None)
                    ]
                    connection.executemany(
                        f"INSERT INTO Produtos ({columns_sql}) VALUES ({placeholders})",
                        records,
                    )

                connection.execute("""
                    CREATE TABLE IF NOT EXISTS VillaColProdutosFonte (
                        Sku TEXT PRIMARY KEY,
                        ReferenciaPdf TEXT NOT NULL,
                        Derivacao TEXT NOT NULL,
                        Descricao TEXT NOT NULL,
                        Categoria TEXT NOT NULL,
                        SubGrupo TEXT NOT NULL,
                        Cor TEXT NOT NULL,
                        PesoKg REAL NOT NULL,
                        Unidade TEXT NOT NULL,
                        QtdeEmbalagemCompra INTEGER NOT NULL,
                        QtdePalete INTEGER NOT NULL,
                        PrecoRevenda REAL NOT NULL,
                        ClassePreco TEXT NOT NULL,
                        ArquivoFonte TEXT NOT NULL,
                        ProcessadoEm TEXT NOT NULL
                    )
                """)
                connection.execute("DELETE FROM VillaColProdutosFonte")
                timestamp = datetime.now().isoformat(timespec="seconds")
                connection.executemany("""
                    INSERT INTO VillaColProdutosFonte (
                        Sku, ReferenciaPdf, Derivacao, Descricao, Categoria,
                        SubGrupo, Cor, PesoKg, Unidade, QtdeEmbalagemCompra,
                        QtdePalete, PrecoRevenda, ClassePreco, ArquivoFonte,
                        ProcessadoEm
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, [
                    (
                        product.sku, product.pdf_reference, product.derivation,
                        product.description, product.category, product.subgroup,
                        product.color, float(product.weight_kg), product.unit,
                        product.package_quantity, product.pallet_quantity,
                        float(product.resale_price), product.price_class,
                        source_pdf.name, timestamp,
                    )
                    for product in products
                ])
                connection.commit()
            except Exception:
                connection.rollback()
                raise

            counts = {
                origin: int(connection.execute(
                    "SELECT COUNT(*) FROM Produtos WHERE TabelaOrigem = ?", (origin,)
                ).fetchone()[0])
                for origin in self.ORIGINS
            }
            counts["VillaColProdutosFonte"] = int(connection.execute(
                "SELECT COUNT(*) FROM VillaColProdutosFonte"
            ).fetchone()[0])
        return counts

    @staticmethod
    def _validate_product_schema(connection: sqlite3.Connection) -> None:
        columns = {
            row[1] for row in connection.execute("PRAGMA table_info(Produtos)").fetchall()
        }
        required = set(DB_COLUMNS) | {"TabelaOrigem"}
        missing = sorted(required - columns)
        if missing:
            raise ValueError(f"Colunas ausentes na tabela Produtos: {missing}")

    @staticmethod
    def _sqlite_value(value: Any) -> Any:
        if value is None or clean_text(value) == "":
            return "" if value == "" else None
        if isinstance(value, Decimal):
            return float(value)
        if hasattr(value, "item"):
            return value.item()
        return value


def write_source_csv(products: Sequence[SourceProduct], output_path: Path) -> None:
    rows = []
    for product in products:
        row = asdict(product)
        row["weight_kg"] = float(product.weight_kg)
        row["resale_price"] = float(product.resale_price)
        rows.append(row)
    pd.DataFrame(rows).to_csv(output_path, index=False, encoding="utf-8-sig")


def write_report(
    output_path: Path,
    audit: AuditResult,
    output_files: Sequence[Path],
    database_counts: Mapping[str, int],
) -> None:
    report = {
        "status": "SUCESSO",
        "processado_em": datetime.now().isoformat(timespec="seconds"),
        "auditoria": asdict(audit),
        "qa": {
            "itens_extraidos_pdf": audit.pdf_sales_skus,
            "itens_planilha_revenda": EXPECTED_SALES_SKUS,
            "totais_batem": audit.pdf_sales_skus == EXPECTED_SALES_SKUS,
            "itens_historicos_por_planilha_del": EXPECTED_DEL_ITEMS,
            "preco_venda_zero": True,
            "codigos_fabrica_unicos": True,
            "arquivos_del_corrigidos": len(EXPECTED_DEL_LEVELS),
        },
        "banco": dict(database_counts),
        "arquivos": [str(path) for path in output_files],
    }
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    project_root = Path(__file__).resolve().parents[1]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--pdf", type=Path, default=None)
    parser.add_argument("--database", type=Path, default=project_root / "R3IntegradorDb.db")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=project_root / "Saida" / f"VILLACOL_PROCESSADO_{timestamp}",
    )
    return parser.parse_args(argv)


def run(args: argparse.Namespace) -> dict[str, Any]:
    input_dir = args.input_dir.resolve()
    pdf_path = (args.pdf or input_dir / DEFAULT_PDF_NAME).resolve()
    database_path = args.database.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    LOGGER.info("Extraindo PDF: %s", pdf_path)
    products = VillaColPdfExtractor().extract(pdf_path)
    LOGGER.info(
        "PDF: %d referências e %d SKUs comerciais.",
        len({product.pdf_reference for product in products}),
        len(products),
    )

    approved = ApprovedSheets(input_dir)
    audit = approved.audit(products)
    if not audit.del30_prices_consistent:
        raise ValueError(
            "DEL30 inconsistente: dispersão da base implícita acima de R$ 0,01."
        )

    writer = ExcelWriter()
    frames: dict[str, pd.DataFrame] = {}
    output_files: list[Path] = []
    for level in EXPECTED_DEL_LEVELS:
        frame = approved.build_del_frame(level)
        origin = f"VILLACOL_DEL{level}"
        frames[origin] = frame
        output_path = output_dir / f"IMPORTACAO_ERP_DELCREDERE_DEL{level}_VILLACOL_CORRIGIDA.xlsx"
        writer.write(approved.paths[level], frame, output_path)
        output_files.append(output_path)

    resale_frame = approved.build_resale_frame(products)
    frames["VILLACOL_REVENDA"] = resale_frame
    resale_path = output_dir / "IMPORTACAO_ERP_REVENDA_VILLACOL.xlsx"
    # DEL20 é usada apenas como molde visual; a fiscalidade já foi consolidada.
    writer.write(approved.paths[20], resale_frame, resale_path)
    output_files.append(resale_path)

    source_csv = output_dir / "ITENS_EXTRAIDOS_PDF_VILLACOL.csv"
    write_source_csv(products, source_csv)
    output_files.append(source_csv)

    database_counts = DatabaseLoader().load(
        database_path, output_dir, frames, products, pdf_path
    )
    expected_database_counts = {
        **{f"VILLACOL_DEL{level}": EXPECTED_DEL_ITEMS for level in EXPECTED_DEL_LEVELS},
        "VILLACOL_REVENDA": EXPECTED_SALES_SKUS,
        "VillaColProdutosFonte": EXPECTED_SALES_SKUS,
    }
    for table, expected_count in expected_database_counts.items():
        count = database_counts[table]
        if count != expected_count:
            raise ValueError(
                f"QA do banco falhou: {table} contém {count}; esperado: {expected_count}."
            )

    report_path = output_dir / "RELATORIO_AUDITORIA_VILLACOL.json"
    write_report(report_path, audit, output_files, database_counts)
    output_files.append(report_path)

    LOGGER.info("[QA OK] Itens extraídos do PDF: %d", len(products))
    LOGGER.info("[QA OK] Itens na importação de revenda: %d", len(resale_frame))
    LOGGER.info("[QA OK] Itens históricos em cada tabela Del Credere: %d", EXPECTED_DEL_ITEMS)
    LOGGER.info("[QA OK] Totais batem perfeitamente: SIM")
    LOGGER.info("[QA OK] DEL30 consistente: SIM (dispersão máxima R$ %.4f)", audit.del30_max_implied_base_spread)
    LOGGER.info("[QA OK] PREÇO VENDA igual a zero em todos os arquivos: SIM")
    product_rows = sum(
        count for table, count in database_counts.items() if table.startswith("VILLACOL_")
    )
    LOGGER.info("[QA OK] Registros ERP persistidos no SQLite: %d", product_rows)
    LOGGER.info(
        "[QA OK] Metadados da fonte persistidos no SQLite: %d",
        database_counts["VillaColProdutosFonte"],
    )
    LOGGER.info("Saída: %s", output_dir)
    return {
        "output_dir": output_dir,
        "files": output_files,
        "audit": audit,
        "database_counts": database_counts,
    }


def main(argv: Sequence[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    try:
        run(parse_args(argv))
        return 0
    except Exception as exc:
        LOGGER.exception("Pipeline VillaCol falhou: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
