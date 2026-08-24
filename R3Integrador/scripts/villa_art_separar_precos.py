#!/usr/bin/env python3
"""Separa a base Villa Art corrigida em atualização e inclusão."""

from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def ref(valor: object) -> str:
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return "" if valor is None else str(valor).strip()


def sha256(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def ler_fonte(caminho: Path) -> dict[str, dict[str, Decimal]]:
    ws = load_workbook(caminho, data_only=True).active
    itens: dict[str, dict[str, Decimal]] = {}
    for linha in ws.iter_rows(min_row=3, values_only=True):
        referencia = ref(linha[1])
        if not referencia or not referencia.isdigit():
            continue
        itens[referencia] = {
            "tabela": Decimal(str(linha[17])),
            # O leitor C# usa o valor formatado a duas casas exibido na tabela.
            "desconto": Decimal(str(linha[18])).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP),
            "na_ponta": Decimal(str(linha[19])),
        }
    return itens


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fonte", required=True, type=Path)
    parser.add_argument("--base-completa", required=True, type=Path)
    parser.add_argument("--banco", required=True, type=Path)
    parser.add_argument("--saida", type=Path)
    args = parser.parse_args()

    saida = args.saida or Path("Saida") / f"VILLA_ART_RECALCULADO_{datetime.now():%Y%m%d_%H%M%S}"
    saida.mkdir(parents=True, exist_ok=False)

    fonte = ler_fonte(args.fonte)
    with sqlite3.connect(args.banco) as conexao:
        cadastradas = {
            ref(linha[0])
            for linha in conexao.execute("SELECT CodigoFabrica FROM Produtos")
            if ref(linha[0])
        }

    wb_base = load_workbook(args.base_completa)
    ws_base = wb_base.active
    ws_base_valores = load_workbook(args.base_completa, data_only=True).active
    linhas = {
        ref(ws_base.cell(numero, 2).value): numero
        for numero in range(2, ws_base.max_row + 1)
        if ref(ws_base.cell(numero, 2).value)
    }
    if set(linhas) != set(fonte):
        raise ValueError("As referências da base completa não coincidem com a fonte.")

    existentes = [referencia for referencia in linhas if referencia in cadastradas]
    novas = [referencia for referencia in linhas if referencia not in cadastradas]

    caminho_atualizacao = saida / "ATUALIZACAO_VILLA_ART_PRECOS_RECALCULADOS.xlsx"
    wb_atualizacao = Workbook()
    ws_atualizacao = wb_atualizacao.active
    ws_atualizacao.title = "ATUALIZACAO"
    ws_atualizacao.append([
        "MARCA", "REFERENCIA", "PREÇO VENDA", "PREÇO DE FÁBRICA",
        "DESCRICAO", "UNID FABRIL", "MODELO", "COR",
    ])
    for celula in ws_atualizacao[1]:
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(horizontal="center")
    for referencia in existentes:
        numero = linhas[referencia]
        ws_atualizacao.append([
            ws_base.cell(numero, 8).value,
            referencia,
            ws_base.cell(numero, 15).value,
            ws_base.cell(numero, 16).value,
            ws_base.cell(numero, 4).value,
            ws_base.cell(numero, 40).value,
            ws_base.cell(numero, 10).value,
            ws_base.cell(numero, 12).value,
        ])
    ws_atualizacao.freeze_panes = "A2"
    ws_atualizacao.auto_filter.ref = ws_atualizacao.dimensions
    for coluna in ("C", "D"):
        for celula in ws_atualizacao[coluna][1:]:
            celula.number_format = 'R$ #,##0.00'
    for coluna, largura in {
        "A": 16, "B": 18, "C": 18, "D": 20,
        "E": 62, "F": 14, "G": 18, "H": 18,
    }.items():
        ws_atualizacao.column_dimensions[coluna].width = largura
    wb_atualizacao.save(caminho_atualizacao)

    caminho_inclusao: Path | None = None
    if novas:
        caminho_inclusao = saida / "IMPORTACAO_ERP_INCLUSAO_VILLA_ART_PRECOS_RECALCULADOS.xlsx"
        for numero in range(ws_base.max_row, 1, -1):
            if ref(ws_base.cell(numero, 2).value) not in novas:
                ws_base.delete_rows(numero)
        wb_base.save(caminho_inclusao)

    itens_auditoria = []
    for referencia, numero in linhas.items():
        preco_final = Decimal(str(ws_base_valores.cell(numero, 15).value))
        esperado = (fonte[referencia]["desconto"] * Decimal("1.1051")).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP)
        if preco_final != esperado:
            raise ValueError(f"Preço divergente para {referencia}: {preco_final} != {esperado}")
        itens_auditoria.append({
            "referencia": referencia,
            "destino": "ATUALIZACAO" if referencia in existentes else "INCLUSAO",
            "preco_tabela": float(fonte[referencia]["tabela"]),
            "preco_desconto": float(fonte[referencia]["desconto"]),
            "preco_sugerido_antigo": float(fonte[referencia]["na_ponta"]),
            "preco_final": float(preco_final),
        })

    arquivos = {caminho_atualizacao.name: sha256(caminho_atualizacao)}
    if caminho_inclusao:
        arquivos[caminho_inclusao.name] = sha256(caminho_inclusao)
    relatorio = {
        "regra": "DESCONTO * (1 + 0.65% IPI + 9.86% ST), sem margem; venda = fábrica",
        "quantidade_fonte": len(fonte),
        "quantidade_atualizacao": len(existentes),
        "quantidade_inclusao": len(novas),
        "referencias_inclusao": novas,
        "arquivos": arquivos,
        "itens": itens_auditoria,
    }
    (saida / "RELATORIO_AUDITORIA_PRECOS.json").write_text(
        json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")
    print(saida.resolve())


if __name__ == "__main__":
    main()
