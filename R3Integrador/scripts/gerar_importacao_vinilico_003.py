#!/usr/bin/env python3
"""Atualiza a importação de vinílicos da tabela 003 com as regras homologadas."""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/home/fernando/Projetos/Work/ARCHOME/ConverterFaltantes/003 TABELA VINÍLICO SP 01.05.26 - ARC HOME (3).xlsx")
TEMPLATE = Path("/home/fernando/Projetos/Work/ARCHOME/Planilhas/IMPORTACAO_ERP_VINILICO_20260701_200706_CONTADOR.xlsx")
OUTPUT = ROOT / "Saida" / "IMPORTACAO_ERP_VINILICO_20260909.xlsx"
REPORT = ROOT / "Saida" / "RELATORIO_AUDITORIA_VINILICO_20260909.json"


def ref(value: object) -> str:
    return "" if value is None else str(value).replace("***", "").strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def final_price(discount: object) -> Decimal:
    base = Decimal(str(discount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return (base * Decimal("1.0792")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def main() -> None:
    source_ws = load_workbook(SOURCE, read_only=True, data_only=True).active
    source: dict[str, dict[str, object]] = {}
    last_m2_box: object = None
    last_weight_box: object = None
    for row in source_ws.iter_rows(min_row=6, max_row=21, values_only=True):
        code = ref(row[1])
        if not code:
            continue
        if row[10] is not None:
            last_m2_box = row[10]
        if row[15] is not None:
            last_weight_box = row[15]
        if last_m2_box is None or last_weight_box is None:
            raise ValueError(f"Embalagem ou peso ausente em {code}")
        source[code] = {
            "preco_final": final_price(row[18]),
            "m2_caixa": last_m2_box,
            "peso_bruto_caixa": last_weight_box,
        }

    book = load_workbook(TEMPLATE)
    ws = book.active
    template_refs = [ref(ws.cell(row, 2).value) for row in range(2, ws.max_row + 1)]
    replacement = {f"LVT9635000{i}": f"LVT9435000{i}" for i in range(1, 5)}
    current_refs = [replacement.get(code, code) for code in template_refs]
    if set(current_refs) != set(source):
        raise ValueError("Referências do perfil do contador não coincidem com a tabela 003 atual")

    for row in range(2, ws.max_row + 1):
        code = replacement.get(ref(ws.cell(row, 2).value), ref(ws.cell(row, 2).value))
        item = source[code]
        ws.cell(row, 2).value = code
        # Regra operacional apontada pela Mônica: preço já líquido nos dois campos
        # e desconto zero, para o ERP não aplicar desconto uma segunda vez.
        ws.cell(row, 15).value = float(item["preco_final"])
        ws.cell(row, 16).value = float(item["preco_final"])
        ws.cell(row, 17).value = 0
        # Peso bruto é o da caixa; embalagem de venda é o m² da caixa da origem.
        ws.cell(row, 25).value = item["m2_caixa"]
        ws.cell(row, 34).value = item["peso_bruto_caixa"]
        for column in (2, 13, 26, 27, 28, 29, 30, 31, 32, 40, 51, 52, 53, 57, 58, 59):
            ws.cell(row, column).number_format = "@"

    ws.auto_filter.ref = ws.dimensions
    book.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=True, data_only=False).active
    rows = list(check.iter_rows(min_row=2, values_only=True))
    mandatory = (1, 3, 5, 7, 12, 13, 14, 15, 18, 19, 23, 25, 26, 27, 28, 29, 30, 31, 38, 39, 50, 56, 57, 58)
    if (check.max_column != 60 or len(rows) != 16 or len({row[1] for row in rows}) != 16 or
            any(not row[index] for row in rows for index in mandatory) or
            any(row[14] != row[15] or row[16] != 0 for row in rows) or
            any(not row[24] or row[24] <= 0 for row in rows) or
            any(cell.data_type == "f" for row in check.iter_rows() for cell in row)):
        raise ValueError("Falha na validação da importação de vinílicos")

    REPORT.write_text(json.dumps({
        "status": "APROVADO_PARA_TESTE_DE_IMPORTACAO",
        "fonte": str(SOURCE), "fonte_sha256": sha256(SOURCE),
        "arquivo": str(OUTPUT), "arquivo_sha256": sha256(OUTPUT),
        "quantidade_saida": len(rows), "colunas": check.max_column,
        "marca": "VILLAGRES", "grupo": "VINILICO", "formulas": 0,
        "regra_preco": "DESCONTO_EXIBIDO_COM_DUAS_CASAS * (1 + 0.0792_ST); IPI_0; SEM_MARKUP",
        "preco_venda_igual_preco_fabrica": True, "desconto_percentual_zerado": True,
        "peso_bruto_origem": "PESO_BRUTO_CAIXA", "embalagem_venda_origem": "M2_POR_CAIXA",
        "perfil_fiscal": "CONTADOR_VINILICO_20260701", "sqlite_alterado": False,
        "bloqueio_rodape": "13 rodapés não gerados para importação: origem sem NCM, unidade e parametrização fiscal homologada.",
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    main()
