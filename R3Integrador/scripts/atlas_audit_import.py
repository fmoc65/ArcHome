#!/usr/bin/env python3
"""Audita a tabela Atlas contra a importação aprovada e carrega o SQLite.

O arquivo aprovado pelo contador é a fonte da parametrização fiscal. A tabela
de preços Atlas é usada para comprovar cobertura e conferir referência, EAN,
descrição, NCM, preço, embalagem, peso e IPI. Os originais não são alterados.
"""

from __future__ import annotations

import argparse
import json
import logging
import math
import os
import re
import shutil
import sqlite3
import sys
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

import pandas as pd
from openpyxl import load_workbook

from villacol_pipeline import DB_COLUMNS, clean_text, decimal_value


LOGGER = logging.getLogger("atlas")
SHEET_SOURCE = "Plan1"
SHEET_ERP = "IMPORTACAO_ERP"
TABLE_ORIGIN = "ATLAS_REVENDA_35"
EXPECTED_PRODUCTS = 203

DEFAULT_DIR = Path("/home/fernando/Projetos/Work/ARCHOME/Planilhas/Atlas")
DEFAULT_SOURCE = "TABELA ATLAS REVENDA 35% - MAIO 2026.xlsx"
DEFAULT_APPROVED = "IMPORTACAO_ERP_ATLAS_REVENDA_35_PROVISORIA_CONTADOR_20260717_OK_CONTADOR.xlsx"


def normalized(value: Any) -> str:
    return clean_text(value).upper()


def number(value: Any) -> float:
    if value is None or clean_text(value) == "":
        return 0.0
    try:
        return float(decimal_value(value))
    except ValueError:
        return 0.0


def ncm_header(text: str) -> tuple[str, str] | None:
    match = re.search(r"NCM\s*-?\s*(\d{4})\.(\d{2})\.(\d{2})", text, re.I)
    if not match:
        return None
    section = re.sub(r"\s*-\s*$", "", text[:match.start()]).strip()
    return section, "".join(match.groups())


def percentage(text: str, name: str) -> float:
    match = re.search(rf"{name}\s*(\d+[,.]?\d*)%", text, re.I)
    return float(match.group(1).replace(",", ".")) if match else 0.0


def extract_source(source_path: Path) -> tuple[list[dict[str, Any]], int]:
    values_book = load_workbook(source_path, data_only=True)
    formulas_book = load_workbook(source_path, data_only=False)
    worksheet = values_book[SHEET_SOURCE]
    formula_sheet = formulas_book[SHEET_SOURCE]
    formula_count = sum(
        cell.data_type == "f"
        for row in formula_sheet.iter_rows()
        for cell in row
    )

    products: list[dict[str, Any]] = []
    section = ""
    ncm = ""
    ipi = 0.0
    source_st = 0.0

    for row_number in range(12, worksheet.max_row + 1):
        column_a = normalized(worksheet.cell(row_number, 1).value)
        header = ncm_header(column_a)
        if header:
            section, ncm = header
            continue
        if "DESCONTO" in column_a:
            ipi = percentage(column_a, "IPI")
            source_st = 0.0 if "ST ISENTO" in column_a else percentage(column_a, "ST")
            continue

        reference = clean_text(worksheet.cell(row_number, 2).value)
        ean = clean_text(worksheet.cell(row_number, 12).value)
        price = number(worksheet.cell(row_number, 5).value)
        if reference and ean and price > 0 and len(ncm) == 8:
            color = normalized(worksheet.cell(row_number, 3).value)
            format_value = normalized(worksheet.cell(row_number, 4).value)
            products.append({
                "source_row": row_number,
                "reference": reference,
                "ean": ean,
                "description": f"{section} - {color} - {format_value}",
                "subgroup": section,
                "model": format_value,
                "color": color,
                "ncm": ncm,
                "price": price,
                "unit": "M2",
                "source_package_sale": number(worksheet.cell(row_number, 9).value) or 1.0,
                "corrected_package_sale": 0.0,
                "gross_weight": number(worksheet.cell(row_number, 10).value),
                "ipi": ipi,
                "source_st": source_st,
            })
            continue

        special_reference = clean_text(worksheet.cell(row_number, 3).value)
        if (
            special_reference
            and special_reference.upper() != "TIPO"
            and price > 0
            and ncm == "69074000"
        ):
            format_value = normalized(worksheet.cell(row_number, 4).value)
            colors = normalized(worksheet.cell(row_number, 9).value)
            if not colors:
                colors = normalized(worksheet.cell(row_number, 10).value)
            products.append({
                "source_row": row_number,
                "reference": special_reference,
                "ean": "",
                "description": f"{section} - {special_reference} - {format_value}",
                "subgroup": section,
                "model": format_value,
                "color": colors,
                "ncm": ncm,
                "price": price,
                "unit": "PC",
                "source_package_sale": 1.0,
                "corrected_package_sale": 1.0,
                "gross_weight": number(worksheet.cell(row_number, 8).value),
                "ipi": ipi,
                "source_st": source_st,
            })

    if len(products) != EXPECTED_PRODUCTS:
        raise ValueError(
            f"Tabela Atlas contém {len(products)} produtos válidos; esperado: {EXPECTED_PRODUCTS}."
        )
    return products, formula_count


def load_approved(approved_path: Path) -> pd.DataFrame:
    frame = pd.read_excel(
        approved_path,
        sheet_name=SHEET_ERP,
        dtype=object,
        keep_default_na=False,
    )
    if len(frame.columns) != 60:
        raise ValueError(f"Layout ERP possui {len(frame.columns)} colunas; esperado: 60.")
    if len(frame) != EXPECTED_PRODUCTS:
        raise ValueError(
            f"Importação aprovada contém {len(frame)} produtos; esperado: {EXPECTED_PRODUCTS}."
        )
    return frame


def compare_source(products: Sequence[dict[str, Any]], frame: pd.DataFrame) -> list[dict[str, Any]]:
    mappings = {
        "reference": "CÓDIGO FÁBRICA",
        "ean": "CODIGOBARRAS",
        "description": "DESCRIÇÃO COMPLETA",
        "subgroup": "SUBGRUPO",
        "model": "MODELO",
        "color": "COR",
        "ncm": "NCM",
        "price": "PREÇO DE FÁBRICA",
        "unit": "UNIDADE",
        "source_package_sale": "QTDE EMBALAGEM DE VENDA",
        "gross_weight": "PESOBRUTO",
        "ipi": "IPI %",
    }
    numeric_fields = {"price", "source_package_sale", "gross_weight", "ipi"}
    mismatches: list[dict[str, Any]] = []
    for index, (source, (_, imported)) in enumerate(zip(products, frame.iterrows()), start=2):
        for source_field, erp_column in mappings.items():
            source_value = source[source_field]
            imported_value = imported[erp_column]
            if source_field in numeric_fields:
                equal = math.isclose(number(source_value), number(imported_value), abs_tol=1e-8)
            else:
                equal = normalized(source_value) == normalized(imported_value)
            if not equal:
                mismatches.append({
                    "erp_row": index,
                    "source_row": source["source_row"],
                    "reference": source["reference"],
                    "field": erp_column,
                    "source": source_value,
                    "approved": imported_value,
                })
    return mismatches


def validate_approved(frame: pd.DataFrame) -> dict[str, Any]:
    required = (
        "CÓDIGO FÁBRICA", "DESCRIÇÃO COMPLETA", "GRUPO", "SUBGRUPO", "MARCA",
        "NCM", "UF ORIGEM", "ALIQICMSORIGEM", "ALIQICMSINTERNA", "CST",
        "ALIQUOTA COFINS CST", "ALIQUOTA IPI CST", "ALIQUOTA PIS CST", "CSOSN",
        "CFOP DENTRO", "CFOP FORA", "ENQUADRAMENTO IPI", "ALIQUOTA IBS",
        "ALIQUOTA CBS", "CLASSIFICACAO TRIBUTARIA",
    )
    missing = {
        column: int(sum(not clean_text(value) for value in frame[column]))
        for column in required
    }
    missing = {column: count for column, count in missing.items() if count}
    if missing:
        raise ValueError(f"Campos obrigatórios vazios: {missing}")

    observations = int(sum(bool(clean_text(value)) for value in frame["OBSERVACAO"]))
    price_mismatches = int(sum(
        not math.isclose(number(sale), number(factory), abs_tol=1e-8)
        for sale, factory in zip(frame["PREÇO VENDA"], frame["PREÇO DE FÁBRICA"])
    ))
    return {
        "missing_required": missing,
        "nonempty_observations": observations,
        "sale_factory_price_mismatches": price_mismatches,
        "unique_references": int(frame["CÓDIGO FÁBRICA"].nunique()),
        "unique_eans_nonempty": int(frame.loc[frame["CODIGOBARRAS"].map(clean_text) != "", "CODIGOBARRAS"].nunique()),
    }


def write_corrected(
    approved_path: Path,
    output_path: Path,
    products: Sequence[dict[str, Any]],
) -> None:
    workbook = load_workbook(approved_path)
    worksheet = workbook[SHEET_ERP]
    for output_row, product in enumerate(products, start=2):
        worksheet.cell(output_row, 25).value = product["corrected_package_sale"]
        worksheet.cell(output_row, 41).value = None
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    os.replace(temporary, output_path)


def sqlite_value(value: Any) -> Any:
    if value is None:
        return None
    if clean_text(value) == "":
        return ""
    if hasattr(value, "item"):
        return value.item()
    return value


def load_database(database_path: Path, output_dir: Path, frame: pd.DataFrame) -> tuple[Path, int]:
    if not database_path.exists():
        raise FileNotFoundError(database_path)
    backup = output_dir / f"{database_path.stem}_ANTES_ATLAS{database_path.suffix}.bak"
    if not backup.exists():
        shutil.copy2(database_path, backup)

    with sqlite3.connect(database_path) as connection:
        columns = {row[1] for row in connection.execute("PRAGMA table_info(Produtos)")}
        required = set(DB_COLUMNS) | {"TabelaOrigem"}
        if required - columns:
            raise ValueError(f"Schema Produtos incompleto: {sorted(required - columns)}")
        connection.execute("BEGIN IMMEDIATE")
        try:
            connection.execute("DELETE FROM Produtos WHERE TabelaOrigem = ?", (TABLE_ORIGIN,))
            all_columns = (*DB_COLUMNS, "TabelaOrigem")
            columns_sql = ",".join(f'"{column}"' for column in all_columns)
            placeholders = ",".join("?" for _ in all_columns)
            records = [
                tuple(sqlite_value(value) for value in row) + (TABLE_ORIGIN,)
                for row in frame.itertuples(index=False, name=None)
            ]
            connection.executemany(
                f"INSERT INTO Produtos ({columns_sql}) VALUES ({placeholders})",
                records,
            )
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        count = int(connection.execute(
            "SELECT COUNT(*) FROM Produtos WHERE TabelaOrigem = ?", (TABLE_ORIGIN,)
        ).fetchone()[0])
    if count != EXPECTED_PRODUCTS:
        raise ValueError(f"SQLite contém {count} produtos Atlas; esperado: {EXPECTED_PRODUCTS}.")
    return backup, count


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_DIR / DEFAULT_SOURCE)
    parser.add_argument("--approved", type=Path, default=DEFAULT_DIR / DEFAULT_APPROVED)
    parser.add_argument("--database", type=Path, default=project_root / "R3IntegradorDb.db")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=project_root / "Saida" / f"ATLAS_PROCESSADO_{datetime.now():%Y%m%d_%H%M%S}",
    )
    return parser.parse_args(argv)


def run(args: argparse.Namespace) -> dict[str, Any]:
    source_path = args.source.resolve()
    approved_path = args.approved.resolve()
    database_path = args.database.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    products, formula_count = extract_source(source_path)
    approved = load_approved(approved_path)
    mismatches = compare_source(products, approved)
    approved_qa = validate_approved(approved)
    if mismatches:
        raise ValueError(f"Foram encontradas {len(mismatches)} divergências contra a origem.")
    if approved_qa["nonempty_observations"]:
        raise ValueError("A planilha aprovada contém observações não vazias.")
    if approved_qa["sale_factory_price_mismatches"]:
        raise ValueError("Preço de venda difere do preço publicado/fábrica.")

    output_path = output_dir / "IMPORTACAO_ERP_ATLAS_REVENDA_35_CORRIGIDA.xlsx"
    write_corrected(approved_path, output_path, products)
    corrected = load_approved(output_path)
    corrected_qa = validate_approved(corrected)
    m2 = corrected[corrected["UNIDADE"].map(normalized) == "M2"]
    pc = corrected[corrected["UNIDADE"].map(normalized) == "PC"]
    if any(number(value) != 0 for value in m2["QTDE EMBALAGEM DE VENDA"]):
        raise ValueError("Há item M2 com quantidade de embalagem de venda diferente de zero.")
    if any(number(value) != 1 for value in pc["QTDE EMBALAGEM DE VENDA"]):
        raise ValueError("Há item PC com quantidade de embalagem de venda diferente de um.")

    backup, database_count = load_database(database_path, output_dir, corrected)
    report = {
        "status": "SUCESSO",
        "processado_em": datetime.now().isoformat(timespec="seconds"),
        "arquivos": {
            "tabela_origem": str(source_path),
            "importacao_aprovada": str(approved_path),
            "importacao_corrigida": str(output_path),
            "backup_sqlite": str(backup),
        },
        "auditoria": {
            "produtos_tabela_origem": len(products),
            "produtos_importacao_aprovada": len(approved),
            "divergencias_de_produtos_precos": len(mismatches),
            "formulas_na_origem": formula_count,
            "formulas_de_preco_na_origem": 0,
            "observacoes_nao_vazias": approved_qa["nonempty_observations"],
            "produtos_m2_corrigidos_embalagem_venda_zero": len(m2),
            "produtos_pc_mantidos_embalagem_venda_um": len(pc),
            "precos_revenda_35_reaplicados": False,
            "campos_fiscais": "preservados da planilha aprovada pelo contador",
        },
        "qa": {
            "totais_batem": len(products) == len(corrected) == database_count,
            "precos_batem_com_origem": True,
            "observacao_vazia": corrected_qa["nonempty_observations"] == 0,
            "quantidade_embalagem_m2_zero": True,
            "registros_sqlite": database_count,
        },
    }
    report_path = output_dir / "RELATORIO_AUDITORIA_ATLAS.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    LOGGER.info("[QA OK] Produtos na origem/importação/SQLite: %d/%d/%d", len(products), len(corrected), database_count)
    LOGGER.info("[QA OK] Divergências de produto ou preço: 0")
    LOGGER.info("[QA OK] Observações não vazias: 0")
    LOGGER.info("[CORRIGIDO] Itens M2 com embalagem de venda ajustada para zero: %d", len(m2))
    LOGGER.info("[QA OK] Itens PC mantidos com embalagem de venda igual a um: %d", len(pc))
    LOGGER.info("Saída: %s", output_dir)
    return report


def main(argv: Sequence[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    try:
        run(parse_args(argv))
        return 0
    except Exception as exc:
        LOGGER.exception("Auditoria Atlas falhou: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
