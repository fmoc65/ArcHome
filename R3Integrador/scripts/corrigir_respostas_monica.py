#!/usr/bin/env python3
"""Gera cópias auditadas das correções pontuais solicitadas por Mônica.

As planilhas de origem nunca são sobrescritas. Nina é processada pelo leitor
.NET, pois a atualização da REV04 faz parte da regra de negócio da aplicação.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ARQUIVO_ADAMA = "IMPORTACAO_ERP_ADAMA_PROVISORIA_20260813_221132.xlsx"
ARQUIVO_STUDIO = "IMPORTACAO_ERP_STUDIO_MORANDIN_PROVISORIA_20260813_214402.xlsx"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_is_valid(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        return archive.testzip() is None


def save_adama(source: Path, destination: Path) -> dict[str, object]:
    workbook = load_workbook(source)
    worksheet = workbook["IMPORTACAO_ERP"]

    # A linha 2 é NCM 32082020, produto químico em SP. Sem outra exceção
    # fornecida, aplica-se a alíquota interna padrão paulista de 18%.
    worksheet.cell(2, 19).value = 18
    worksheet.cell(2, 20).value = 18
    worksheet.cell(2, 19).number_format = "0"
    worksheet.cell(2, 20).number_format = "0"
    workbook.save(destination)

    return {
        "arquivo": destination.name,
        "linhas": worksheet.max_row - 1,
        "referencia": worksheet.cell(2, 2).value,
        "ncm": worksheet.cell(2, 13).value,
        "icms_origem": worksheet.cell(2, 19).value,
        "icms_interna": worksheet.cell(2, 20).value,
        "xlsx_integro": workbook_is_valid(destination),
    }


def save_studio(source: Path, destination: Path) -> dict[str, object]:
    workbook = load_workbook(source)
    worksheet = workbook["IMPORTACAO_ERP"]
    updated = 0

    for row in range(2, worksheet.max_row + 1):
        cell = worksheet.cell(row, 30)
        if str(cell.value).strip() == "0500":
            cell.value = "500"
            updated += 1
        cell.number_format = "@"

    workbook.save(destination)
    return {
        "arquivo": destination.name,
        "linhas": worksheet.max_row - 1,
        "csosn_alterados": updated,
        "csosn_distintos": sorted(
            {str(worksheet.cell(row, 30).value) for row in range(2, worksheet.max_row + 1)}
        ),
        "xlsx_integro": workbook_is_valid(destination),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--entrada", type=Path, required=True)
    parser.add_argument("--saida", type=Path, required=True)
    args = parser.parse_args()

    args.saida.mkdir(parents=True, exist_ok=False)
    adama_source = args.entrada / ARQUIVO_ADAMA
    studio_source = args.entrada / ARQUIVO_STUDIO
    if not adama_source.is_file() or not studio_source.is_file():
        raise FileNotFoundError("As planilhas Adama e Studio Morandin não foram localizadas.")

    adama_output = args.saida / "IMPORTACAO_ERP_ADAMA_CORRIGIDA_MONICA_20260907.xlsx"
    studio_output = args.saida / "IMPORTACAO_ERP_STUDIO_MORANDIN_CORRIGIDA_MONICA_20260907.xlsx"
    report = {
        "fontes": {
            adama_source.name: sha256(adama_source),
            studio_source.name: sha256(studio_source),
        },
        "adama": save_adama(adama_source, adama_output),
        "studio_morandin": save_studio(studio_source, studio_output),
        "saidas": {
            adama_output.name: sha256(adama_output),
            studio_output.name: sha256(studio_output),
        },
    }
    (args.saida / "RELATORIO_AUDITORIA_ADAMA_STUDIO_MONICA_20260907.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(args.saida.resolve())


if __name__ == "__main__":
    main()
