#!/usr/bin/env python3
"""Gera e valida o layout ERP completo da tabela Villa Art revenda."""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


RAIZ = Path(__file__).resolve().parents[1]
FONTE = Path("/home/fernando/Projetos/Work/ARCHOME/ConverterFaltantes/002 TABELA VILLA ART SP 01.05.26 - ARC HOME (3).xlsx")
MODELO = RAIZ / "Saida" / "IMPORTACAO_ERP_VILLA_ART_20260901_190836.xlsx"
SAIDA = RAIZ / "Saida" / "IMPORTACAO_ERP_VILLA_ART_20260909.xlsx"
RELATORIO = RAIZ / "Saida" / "RELATORIO_AUDITORIA_VILLA_ART_20260909.json"


def txt(value: object) -> str:
    return "" if value is None else str(value).strip()


def money(value: object) -> Decimal:
    discount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return (discount * Decimal("1.1051")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    source = load_workbook(FONTE, data_only=True, read_only=True).active
    book = load_workbook(MODELO)
    target = book.active
    target.delete_rows(2, target.max_row - 1)

    products = []
    last_format = ""
    last_m2_box: object = None
    last_weight_box: object = None
    for row in source.iter_rows(min_row=3, values_only=True):
        reference = txt(row[1])
        if not reference:
            continue
        if txt(row[0]):
            last_format = txt(row[0]).replace(" ", "").upper()
        if row[10] is not None:
            last_m2_box = row[10]
        if row[15] is not None:
            last_weight_box = row[15]
        if not last_format:
            raise ValueError(f"Formato ausente: {reference}")
        if last_m2_box is None or last_weight_box is None:
            raise ValueError(f"Embalagem ou peso ausente: {reference}")
        line, surface, color = txt(row[2]).upper(), txt(row[5]).upper(), txt(row[4]).lower()
        price = money(row[18])
        values = [
            "", reference, "", f"PORCELANATO {line} {surface} {color.upper()} {last_format}",
            f"PORCELANATO {last_format} {color.upper()}", "PORCELANATO", surface,
            "VILLA ART", line, last_format, "", color, "69072100", "SP",
            float(price), float(price), 0, 0.65, 12, 12, 81, 0, 0, "M2", last_m2_box,
            "010", "01", "49", "01", "500", "5405", "6404", 0, last_weight_box, 1,
            0, 0, 0, 9.86, "CX", "", 0, 0, 0, "", "", "", "", "", "",
            "999", "0,65", "3", "", 0, 0, "0,1", "0,9", "000001", "",
        ]
        if len(values) != 60:
            raise ValueError("Layout diferente de 60 colunas")
        target.append(values)
        products.append({"referencia": reference, "preco_final": float(price), "m2_caixa": last_m2_box, "peso_bruto_caixa": last_weight_box})

    target.auto_filter.ref = target.dimensions
    book.save(SAIDA)

    check = load_workbook(SAIDA, data_only=False, read_only=True).active
    rows = list(check.iter_rows(min_row=2, values_only=True))
    fiscal_columns = (12, 13, 17, 18, 19, 20, 25, 38, 50, 56, 57, 58)
    if (len(rows) != len(products) or check.max_column != 60 or
            len({row[1] for row in rows}) != len(rows) or
            any(row[7] != "VILLA ART" for row in rows) or
            any(not row[24] or row[24] <= 0 for row in rows) or
            any(not row[index] for row in rows for index in fiscal_columns) or
            any(row[14] != row[15] or row[16] != 0 for row in rows) or
            any(cell.data_type == "f" for row in check.iter_rows() for cell in row)):
        raise ValueError("Falha na validação da saída ERP")

    RELATORIO.write_text(json.dumps({
        "status": "APROVADO_PARA_TESTE_DE_IMPORTACAO",
        "finalidade": "IMPORTACAO_COMPLETA_VILLA_ART_REVENDA",
        "fonte": str(FONTE), "fonte_sha256": digest(FONTE),
        "arquivo": str(SAIDA), "arquivo_sha256": digest(SAIDA),
        "quantidade_fonte": len(products), "quantidade_saida": len(rows),
        "referencias_unicas": len({row[1] for row in rows}), "colunas": check.max_column,
        "formulas": 0, "marca": "VILLA ART",
        "regra_preco": "PRECO_DESCONTO_EXIBIDO_COM_DUAS_CASAS * (1 + 0.0065 + 0.0986), ARREDONDADO_EM_DUAS_CASAS",
        "preco_venda_igual_preco_fabrica": True, "desconto_percentual_zerado": True,
        "peso_bruto_origem": "PESO_BRUTO_CAIXA", "embalagem_venda_origem": "M2_POR_CAIXA",
        "sqlite_alterado": False, "itens": products,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(SAIDA)


if __name__ == "__main__":
    main()
