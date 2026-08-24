#!/usr/bin/env python3
"""Refaz as planilhas VAREJO usando desconto + IPI + ST, sem markup."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


CENTAVOS = Decimal("0.01")
ST_PERCENTUAL = Decimal("9.86")
IPI_PADRAO = Decimal("0.65")


def decimal(valor: object) -> Decimal:
    return Decimal(str(valor))


def dinheiro(valor: Decimal) -> Decimal:
    return valor.quantize(CENTAVOS, rounding=ROUND_HALF_UP)


def referencia(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def sha256(caminho: Path) -> str:
    digest = hashlib.sha256()
    with caminho.open("rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(1024 * 1024), b""):
            digest.update(bloco)
    return digest.hexdigest()


def ler_precos(fonte: Path) -> dict[str, Decimal]:
    planilha = load_workbook(fonte, data_only=True).active
    precos: dict[str, Decimal] = {}
    for linha in planilha.iter_rows(min_row=3, values_only=True):
        ref = referencia(linha[1])
        preco_desconto = linha[18]
        if not ref or not isinstance(preco_desconto, (int, float, Decimal)):
            continue
        if ref in precos:
            raise ValueError(f"Referência duplicada na fonte: {ref}")
        precos[ref] = decimal(preco_desconto)
    return precos


def ler_linhas(caminho: Path) -> tuple[list[str], dict[str, list[object]]]:
    planilha = load_workbook(caminho, data_only=True).active
    cabecalho = [celula.value for celula in planilha[1]]
    linhas: dict[str, list[object]] = {}
    for linha in planilha.iter_rows(min_row=2, values_only=True):
        ref = referencia(linha[1])
        if not ref:
            continue
        if ref in linhas:
            raise ValueError(f"Referência duplicada em {caminho.name}: {ref}")
        linhas[ref] = list(linha)
    return cabecalho, linhas


def ipi_percentual(ref: str, linha: list[object]) -> Decimal:
    # A regra fiscal vigente do mapper usa IPI 0,65% para NCM 69072100.
    # A referência 120003 é NCM 69072200 e, portanto, tem IPI zero.
    if ref == "120003":
        return Decimal("0")
    ipi_existente = linha[17]
    if isinstance(ipi_existente, (int, float, Decimal)) and decimal(ipi_existente) > 0:
        return decimal(ipi_existente)
    return IPI_PADRAO


def calcular(preco_desconto: Decimal, ipi: Decimal) -> Decimal:
    acrescimo = Decimal("1") + ipi / Decimal("100") + ST_PERCENTUAL / Decimal("100")
    return dinheiro(preco_desconto * acrescimo)


def criar_atualizacao(
    destino: Path,
    linhas: dict[str, list[object]],
    precos: dict[str, Decimal],
) -> list[dict[str, object]]:
    wb = Workbook()
    ws = wb.active
    ws.title = "ATUALIZACAO"
    cabecalho = [
        "MARCA",
        "REFERENCIA",
        "PREÇO VENDA",
        "PREÇO DE FÁBRICA",
        "DESCRICAO",
        "UNID FABRIL",
        "MODELO",
        "COR",
    ]
    ws.append(cabecalho)
    fill = PatternFill("solid", fgColor="1F4E78")
    for celula in ws[1]:
        celula.fill = fill
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(horizontal="center")

    auditoria: list[dict[str, object]] = []
    for ref, linha in linhas.items():
        if ref not in precos:
            raise ValueError(f"Referência da atualização ausente na fonte: {ref}")
        ipi = ipi_percentual(ref, linha)
        novo_preco = calcular(precos[ref], ipi)
        ws.append([
            linha[7],
            ref,
            float(novo_preco),
            float(novo_preco),
            linha[3],
            linha[39],
            linha[9],
            linha[11],
        ])
        auditoria.append({
            "referencia": ref,
            "preco_desconto": float(precos[ref]),
            "ipi_percentual": float(ipi),
            "st_percentual": float(ST_PERCENTUAL),
            "preco_final": float(novo_preco),
        })

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for coluna in ("C", "D"):
        for celula in ws[coluna][1:]:
            celula.number_format = 'R$ #,##0.00'
    larguras = {"A": 16, "B": 18, "C": 18, "D": 20, "E": 55, "F": 14, "G": 20, "H": 20}
    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura
    wb.save(destino)
    return auditoria


def criar_inclusao(
    origem: Path,
    destino: Path,
    precos: dict[str, Decimal],
) -> list[dict[str, object]]:
    wb = load_workbook(origem)
    ws = wb.active
    auditoria: list[dict[str, object]] = []
    for numero_linha in range(2, ws.max_row + 1):
        ref = referencia(ws.cell(numero_linha, 2).value)
        if not ref:
            continue
        if ref not in precos:
            raise ValueError(f"Referência da inclusão ausente na fonte: {ref}")
        valores = [ws.cell(numero_linha, coluna).value for coluna in range(1, ws.max_column + 1)]
        ipi = ipi_percentual(ref, valores)
        novo_preco = calcular(precos[ref], ipi)
        ws.cell(numero_linha, 15).value = float(novo_preco)
        ws.cell(numero_linha, 16).value = float(novo_preco)
        ws.cell(numero_linha, 15).number_format = 'R$ #,##0.00'
        ws.cell(numero_linha, 16).number_format = 'R$ #,##0.00'
        auditoria.append({
            "referencia": ref,
            "preco_desconto": float(precos[ref]),
            "ipi_percentual": float(ipi),
            "st_percentual": float(ST_PERCENTUAL),
            "preco_final": float(novo_preco),
        })
    wb.save(destino)
    return auditoria


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fonte", type=Path, required=True)
    parser.add_argument("--atualizacao-base", type=Path, required=True)
    parser.add_argument("--inclusao-base", type=Path, required=True)
    parser.add_argument("--saida", type=Path)
    args = parser.parse_args()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    saida = args.saida or Path("Saida") / f"VAREJO_RECALCULADO_{timestamp}"
    saida.mkdir(parents=True, exist_ok=False)

    precos = ler_precos(args.fonte)
    _, linhas_atualizacao = ler_linhas(args.atualizacao_base)
    _, linhas_inclusao = ler_linhas(args.inclusao_base)
    intersecao = set(linhas_atualizacao) & set(linhas_inclusao)
    if intersecao:
        raise ValueError(f"Referências presentes nos dois arquivos: {sorted(intersecao)}")
    processadas = set(linhas_atualizacao) | set(linhas_inclusao)
    if processadas != set(precos):
        raise ValueError(
            f"Cobertura divergente: fonte={len(precos)}, processadas={len(processadas)}, "
            f"faltantes={sorted(set(precos) - processadas)}, extras={sorted(processadas - set(precos))}"
        )

    atualizacao = saida / "ATUALIZACAO_VAREJO_PRECOS_RECALCULADOS.xlsx"
    inclusao = saida / "IMPORTACAO_ERP_INCLUSAO_VAREJO_PRECOS_RECALCULADOS_PENDENTE_FISCAL.xlsx"
    audit_atualizacao = criar_atualizacao(atualizacao, linhas_atualizacao, precos)
    audit_inclusao = criar_inclusao(args.inclusao_base, inclusao, precos)

    relatorio = {
        "regra": (
            "preco_final = arredondar(preco_desconto * "
            "(1 + ipi_percentual/100 + 9.86/100), 2); sem markup; "
            "preco_venda = preco_fabrica = preco_final"
        ),
        "fonte": str(args.fonte),
        "quantidade_fonte": len(precos),
        "quantidade_atualizacao": len(audit_atualizacao),
        "quantidade_inclusao": len(audit_inclusao),
        "arquivos": {
            atualizacao.name: sha256(atualizacao),
            inclusao.name: sha256(inclusao),
        },
        "itens_atualizacao": audit_atualizacao,
        "itens_inclusao": audit_inclusao,
    }
    relatorio_path = saida / "RELATORIO_AUDITORIA_PRECOS.json"
    relatorio_path.write_text(json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")
    print(saida.resolve())


if __name__ == "__main__":
    main()
