#!/usr/bin/env python3
"""Gera o layout completo do ERP para a tabela de lançamento HOUSE."""

from __future__ import annotations

import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


RAIZ = Path(__file__).resolve().parents[1]
FONTE = Path("/home/fernando/Downloads/ultimas/TABELA HOUSE LANÇAMENTO 2026.xlsx")
MODELO = RAIZ / "Saida" / "IMPORTACAO_ERP_VILLA_ART_20260901_190836.xlsx"
SAIDA = RAIZ / "Saida" / "IMPORTACAO_ERP_HOUSE_VILLAGRES_20261005.xlsx"
RELATORIO = RAIZ / "Saida" / "RELATORIO_AUDITORIA_HOUSE_VILLAGRES_20261005.json"


def texto(valor: object) -> str:
    return "" if valor is None else str(valor).strip()


def modelo(formato: object) -> str:
    return texto(formato).replace(" ", "").upper()


def preco_final(valor_fracionado: object) -> Decimal:
    # A tabela atual já traz o acréscimo de fracionado na última coluna.
    # Usa esse valor (arredondado a centavos) e soma IPI 0,65% + ST 9,86%,
    # sem reaplicar o adicional de 10%.
    fracionado = Decimal(str(valor_fracionado)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return (fracionado * Decimal("1.1051")).quantize(
        Decimal("0.01"), rounding=ROUND_HALF_UP)


def sha256(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def main() -> None:
    ws_fonte = load_workbook(FONTE, data_only=True, read_only=True).active
    if texto(ws_fonte.cell(2, 20).value).upper() != "FRACIONADO":
        raise ValueError("A coluna T/20 da tabela não está identificada como FRACIONADO")
    wb_saida = load_workbook(MODELO)
    ws_saida = wb_saida.active

    # Mantém cabeçalho, estilos, filtros e larguras do último layout ERP aprovado.
    if ws_saida.max_row > 1:
        ws_saida.delete_rows(2, ws_saida.max_row - 1)

    itens = []
    ultimo_formato = ""
    for linha in ws_fonte.iter_rows(min_row=3, values_only=True):
        referencia = texto(linha[1])
        if not referencia:
            continue

        preco_fracionado = Decimal(str(linha[19])).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        preco = preco_final(preco_fracionado)
        superficie = texto(linha[5]).upper()
        cor = texto(linha[4]).lower()
        # A origem deixa o formato vazio nas linhas subsequentes do mesmo bloco.
        if texto(linha[0]):
            ultimo_formato = modelo(linha[0])
        tamanho = ultimo_formato
        if not tamanho:
            raise ValueError(f"Formato ausente para a referência {referencia}")
        peso_caixa = linha[15]
        registro = [
            "", referencia, "",
            f"PORCELANATO HOUSE {superficie} {cor.upper()} {tamanho}".strip(),
            f"PORCELANATO {tamanho} {cor.upper()}".strip(),
            "PORCELANATO", superficie, "VILLAGRES", "HOUSE", tamanho, "", cor,
            "69072100", "SP", float(preco), float(preco), 0, 0.65, 12, 12, 81,
            0, 0, "M2", linha[10], "010", "01", "49", "01", "500", "5405", "6404",
            0, peso_caixa, 1, 0, 0, 0, 9.86, "CX", "", 0, 0, 0,
            "", "", "", "", "", "", "999", "0,65", "3", "", 0, 0,
            "0,1", "0,9", "000001", "",
        ]
        if len(registro) != 60:
            raise ValueError(f"Layout inválido: {len(registro)} colunas")
        ws_saida.append(registro)
        itens.append({
            "referencia": referencia,
            "preco_fracionado_fonte": float(preco_fracionado),
            "preco_final": float(preco),
            "qtde_embalagem_venda_m2_caixa": linha[10],
            "peso_bruto_caixa": peso_caixa,
        })

    ws_saida.auto_filter.ref = ws_saida.dimensions
    wb_saida.save(SAIDA)

    # QA no arquivo persistido, sem confiar no estado mantido em memória.
    ws_qa = load_workbook(SAIDA, data_only=False, read_only=True).active
    linhas = list(ws_qa.iter_rows(min_row=2, values_only=True))
    if len(linhas) != 11 or ws_qa.max_column != 60:
        raise ValueError("Quantidade de itens ou número de colunas divergente")
    referencias = [linha[1] for linha in linhas]
    if len(set(referencias)) != len(referencias) or any(linha[7] != "VILLAGRES" for linha in linhas):
        raise ValueError("Referências ou marca inválidas")
    if any(linha[14] != linha[15] or linha[16] != 0 for linha in linhas):
        raise ValueError("Preço/desconto inválido")
    if any(not linha[24] or linha[24] <= 0 for linha in linhas):
        raise ValueError("Quantidade de embalagem de venda inválida")
    if any(not linha[indice] for linha in linhas for indice in (12, 13, 50, 56, 57, 58)):
        raise ValueError("Campo fiscal obrigatório ausente")
    if any(celula.data_type == "f" for linha in ws_qa.iter_rows() for celula in linha):
        raise ValueError("A saída não pode conter fórmulas")

    RELATORIO.write_text(json.dumps({
        "status": "APROVADO_PARA_TESTE_DE_IMPORTACAO",
        "finalidade": "IMPORTACAO_COMPLETA_HOUSE_VILLAGRES",
        "fonte": str(FONTE),
        "fonte_sha256": sha256(FONTE),
        "arquivo": str(SAIDA),
        "arquivo_sha256": sha256(SAIDA),
        "quantidade_fonte": len(itens),
        "quantidade_saida": len(linhas),
        "referencias_unicas": len(set(referencias)),
        "colunas": ws_qa.max_column,
        "formulas": 0,
        "marca": "VILLAGRES",
        "regra_preco": "PRECO_FRACIONADO_COLUNA_T_COM_DUAS_CASAS * (1 + 0.0065 + 0.0986), SEM_REAPLICAR_10_PORCENTO",
        "preco_venda_igual_preco_fabrica": True,
        "desconto_percentual_zerado": True,
        "peso_bruto_origem": "PESO_BRUTO_CAIXA",
        "sqlite_alterado": False,
        "itens": itens,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(SAIDA)
    print(RELATORIO)


if __name__ == "__main__":
    main()
