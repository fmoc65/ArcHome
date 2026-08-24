#!/usr/bin/env python3
"""Separa a tabela 003 em vinílicos e rodapés e reconcilia com o SQLite."""

from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def referencia(valor: object) -> str:
    if valor is None:
        return ""
    return str(valor).replace("***", "").strip()


def sha256(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def criar_atualizacao(
    caminho: Path,
    referencias: list[str],
    linhas: dict[str, int],
    ws_base,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ATUALIZACAO"
    ws.append([
        "MARCA", "REFERENCIA", "PREÇO VENDA", "PREÇO DE FÁBRICA",
        "DESCRICAO", "UNID FABRIL", "MODELO", "COR",
    ])
    for celula in ws[1]:
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(horizontal="center")
    for ref in referencias:
        numero = linhas[ref]
        ws.append([
            ws_base.cell(numero, 8).value,
            ref,
            ws_base.cell(numero, 15).value,
            ws_base.cell(numero, 16).value,
            ws_base.cell(numero, 4).value,
            ws_base.cell(numero, 40).value,
            ws_base.cell(numero, 10).value,
            ws_base.cell(numero, 12).value,
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for coluna in ("C", "D"):
        for celula in ws[coluna][1:]:
            celula.number_format = 'R$ #,##0.00'
    for coluna, largura in {
        "A": 16, "B": 18, "C": 18, "D": 20,
        "E": 60, "F": 14, "G": 18, "H": 18,
    }.items():
        ws.column_dimensions[coluna].width = largura
    wb.save(caminho)


def criar_inclusao(base: Path, caminho: Path, referencias: set[str], segmento: str) -> None:
    wb = load_workbook(base)
    ws = wb.active
    for numero in range(ws.max_row, 1, -1):
        if referencia(ws.cell(numero, 2).value) not in referencias:
            ws.delete_rows(numero)
    if segmento == "RODAPE":
        colunas_fiscais_pendentes = [
            13, 18, 19, 20, 21, 26, 27, 28, 29, 30, 31, 32,
            36, 37, 38, 39, 42, 43, 44, 51, 52, 53, 57, 58, 59, 60,
        ]
        for numero in range(2, ws.max_row + 1):
            subgrupo = str(ws.cell(numero, 7).value or "").strip()
            cor = str(ws.cell(numero, 12).value or "").strip()
            modelo = str(ws.cell(numero, 10).value or "").strip()
            ws.cell(numero, 4).value = f"RODAPE POLIESTIRENO {subgrupo} {cor} {modelo}".upper()
            ws.cell(numero, 5).value = f"RODAPE {modelo} {cor}".upper()
            ws.cell(numero, 6).value = "RODAPE"
            ws.cell(numero, 24).value = ""
            ws.cell(numero, 41).value = (
                "Tabela 003 Villagres - rodape separado; unidade, NCM e fiscal pendentes"
            )
            for coluna in colunas_fiscais_pendentes:
                ws.cell(numero, coluna).value = ""
    wb.save(caminho)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--fonte", required=True, type=Path)
    parser.add_argument("--base-completa", required=True, type=Path)
    parser.add_argument("--banco", required=True, type=Path)
    parser.add_argument("--saida", type=Path)
    args = parser.parse_args()

    saida = args.saida or Path("Saida") / f"VINILICO_RODAPE_PROCESSADO_{datetime.now():%Y%m%d_%H%M%S}"
    saida.mkdir(parents=True, exist_ok=False)

    ws_fonte = load_workbook(args.fonte, data_only=True).active
    fonte: dict[str, dict[str, object]] = {}
    for linha in ws_fonte.iter_rows(min_row=1, values_only=True):
        ref = referencia(linha[1] if len(linha) > 1 else None)
        if not (ref.startswith("SPC") or ref.startswith("LVT") or ref.startswith("RP")):
            continue
        desconto = Decimal(str(linha[18])).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        fonte[ref] = {
            "segmento": "RODAPE" if ref.startswith("RP") else "VINILICO",
            "preco_tabela": float(Decimal(str(linha[17]))),
            "preco_desconto": float(desconto),
        }

    wb_base = load_workbook(args.base_completa, data_only=True)
    ws_base = wb_base.active
    linhas = {
        referencia(ws_base.cell(numero, 2).value): numero
        for numero in range(2, ws_base.max_row + 1)
        if referencia(ws_base.cell(numero, 2).value)
    }
    if set(linhas) != set(fonte):
        raise ValueError("As referências da base completa não coincidem com a fonte.")

    with sqlite3.connect(args.banco) as conexao:
        cadastradas = {
            referencia(linha[0])
            for linha in conexao.execute("SELECT CodigoFabrica FROM Produtos")
            if referencia(linha[0])
        }

    arquivos: dict[str, str] = {}
    resumo: dict[str, dict[str, object]] = {}
    itens = []
    for segmento in ("VINILICO", "RODAPE"):
        refs_segmento = [ref for ref in linhas if fonte[ref]["segmento"] == segmento]
        existentes = [ref for ref in refs_segmento if ref in cadastradas]
        novas = [ref for ref in refs_segmento if ref not in cadastradas]
        resumo[segmento] = {
            "fonte": len(refs_segmento),
            "atualizacao": len(existentes),
            "inclusao": len(novas),
        }

        if existentes:
            caminho = saida / f"ATUALIZACAO_{segmento}_PRECOS_RECALCULADOS.xlsx"
            criar_atualizacao(caminho, existentes, linhas, ws_base)
            arquivos[caminho.name] = sha256(caminho)
        if novas:
            sufixo = "_PENDENTE_UNIDADE_E_FISCAL" if segmento == "RODAPE" else ""
            caminho = saida / f"IMPORTACAO_ERP_INCLUSAO_{segmento}{sufixo}.xlsx"
            criar_inclusao(args.base_completa, caminho, set(novas), segmento)
            arquivos[caminho.name] = sha256(caminho)

        for ref in refs_segmento:
            numero = linhas[ref]
            preco_final = Decimal(str(ws_base.cell(numero, 15).value))
            esperado = (Decimal(str(fonte[ref]["preco_desconto"])) * Decimal("1.0792")).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP)
            if preco_final != esperado or ws_base.cell(numero, 16).value != ws_base.cell(numero, 15).value:
                raise ValueError(f"Preço inválido para {ref}.")
            itens.append({
                "referencia": ref,
                "segmento": segmento,
                "destino": "ATUALIZACAO" if ref in cadastradas else "INCLUSAO",
                **fonte[ref],
                "preco_final": float(preco_final),
            })

    relatorio = {
        "regra_preco": "DESCONTO * (1 + 7.92% ST); IPI 0%; sem R$ 1,50; sem markup; venda = fábrica",
        "quantidade_total": len(fonte),
        "resumo": resumo,
        "pendencia_rodape": (
            "A fonte não informa NCM nem tributação específica do rodapé de poliestireno. "
            "O arquivo foi separado, mas unidade e fiscal precisam de homologação antes da carga."
        ),
        "arquivos": arquivos,
        "itens": itens,
    }
    (saida / "RELATORIO_AUDITORIA_PRECOS_E_BANCO.json").write_text(
        json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")
    print(saida.resolve())


if __name__ == "__main__":
    main()
